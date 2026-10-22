"""
Build a visual, low-text capstone deck for the Clinical Report Assistant.

Design follows the supplied MiniMentigo make_slides.py language:
light slides, navy/teal/amber palette, vector boxes and arrows, short labels.
Diagrams carry the story; prose is deliberately sparse.

Run:
    uv run --with python-pptx python scripts/build_capstone_deck_visual.py
Output:
    Clinical_Report_Assistant_Presentation.pptx (project root)
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

# ---------------------------------------------------------------------------
# Palette
# ---------------------------------------------------------------------------
NAVY = RGBColor(0x0F, 0x2A, 0x43)
NAVY_SOFT = RGBColor(0x1B, 0x3D, 0x5C)
TEAL = RGBColor(0x12, 0xB3, 0xA6)
TEAL_DARK = RGBColor(0x0E, 0x8C, 0x82)
TEAL_LIGHT = RGBColor(0xE4, 0xF7, 0xF4)
AMBER = RGBColor(0xF2, 0xA9, 0x3B)
AMBER_DARK = RGBColor(0xB4, 0x79, 0x1E)
AMBER_LIGHT = RGBColor(0xFD, 0xF2, 0xDF)
PLUM = RGBColor(0x6B, 0x5B, 0xA8)
PLUM_LIGHT = RGBColor(0xEF, 0xEC, 0xF7)
LIGHT = RGBColor(0xF5, 0xF7, 0xFA)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
TEXT = RGBColor(0x1B, 0x27, 0x33)
MUTED = RGBColor(0x5B, 0x6B, 0x7B)
BORDER = RGBColor(0xD6, 0xDE, 0xE6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PALE_TEXT = RGBColor(0xC9, 0xD6, 0xDF)

FONT = "Segoe UI"
BULLET = "\u2022  "

SW, SH = 13.333, 7.5
TOTAL = 0


# ---------------------------------------------------------------------------
# Low-level helpers
# ---------------------------------------------------------------------------
def no_line(shape) -> None:
    shape.line.fill.background()


def no_shadow(shape) -> None:
    shape.shadow.inherit = False


def fill(shape, color: RGBColor) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def set_text(tf, lines, *, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             margin=0.08, wrap=True) -> None:
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    for i, spec in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = spec.get("align", align)
        p.space_before = Pt(spec.get("space_before", 0))
        p.space_after = Pt(spec.get("space_after", 2))
        run = p.add_run()
        run.text = spec.get("text", "")
        run.font.size = Pt(spec.get("size", 14))
        run.font.bold = spec.get("bold", False)
        run.font.italic = spec.get("italic", False)
        run.font.name = FONT
        run.font.color.rgb = spec.get("color", TEXT)


def add_textbox(slide, l, t, w, h, lines, **kw):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    set_text(box.text_frame, lines, **kw)
    return box


def add_box(slide, l, t, w, h, lines, *, color=CARD, line=BORDER, radius=0.09,
            line_w=1.0, dash=None, align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h)
    )
    shape.adjustments[0] = radius
    fill(shape, color)
    if line is None:
        no_line(shape)
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(line_w)
        if dash is not None:
            shape.line.dash_style = dash
    no_shadow(shape)
    set_text(shape.text_frame, lines, anchor=anchor, align=align)
    return shape


def add_rect(slide, l, t, w, h, color, *, line=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h)
    )
    fill(shape, color)
    if line is None:
        no_line(shape)
    else:
        shape.line.color.rgb = line
    no_shadow(shape)
    return shape


def add_arrow(slide, x1, y1, x2, y2, *, color=TEAL, width=2.1, dash=None):
    conn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    conn.line.color.rgb = color
    conn.line.width = Pt(width)
    if dash is not None:
        conn.line.dash_style = dash
    ln = conn.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    no_shadow(conn)
    return conn


def add_line(slide, x1, y1, x2, y2, *, color=TEAL, width=2.1, dash=None):
    conn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    conn.line.color.rgb = color
    conn.line.width = Pt(width)
    if dash is not None:
        conn.line.dash_style = dash
    no_shadow(conn)
    return conn


def add_badge(slide, l, t, size, text, color=TEAL, fsize=13):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(l), Inches(t), Inches(size), Inches(size)
    )
    fill(shape, color)
    no_line(shape)
    no_shadow(shape)
    set_text(
        shape.text_frame,
        [{"text": str(text), "size": fsize, "bold": True, "color": WHITE,
          "align": PP_ALIGN.CENTER, "space_after": 0}],
        anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, margin=0.0,
    )
    return shape


def add_pill(slide, l, t, w, h, text, *, color=CARD, line=BORDER,
             text_color=NAVY, size=13, bold=True, radius=0.5):
    return add_box(
        slide, l, t, w, h,
        [{"text": text, "size": size, "bold": bold, "color": text_color,
          "space_after": 0}],
        color=color, line=line, radius=radius,
    )


def add_footer(slide, page: int) -> None:
    add_textbox(slide, 0.55, 7.05, 9.5, 0.3,
                [{"text": "AI Clinical Report Summarization Assistant  \u00b7  "
                          "synthetic-data capstone",
                  "size": 10, "color": MUTED, "space_after": 0}])
    add_textbox(slide, 11.4, 7.05, 1.4, 0.3,
                [{"text": f"{page} / {TOTAL}", "size": 10, "color": MUTED,
                  "align": PP_ALIGN.RIGHT, "space_after": 0}])


def new_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def add_header(slide, title, subtitle=None):
    add_textbox(slide, 0.6, 0.38, 12.1, 0.72,
                [{"text": title, "size": 27, "bold": True, "color": NAVY,
                  "space_after": 0}])
    bar = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.63), Inches(1.13), Inches(1.5), Inches(0.07)
    )
    fill(bar, TEAL)
    no_line(bar)
    no_shadow(bar)
    if subtitle:
        add_textbox(slide, 0.6, 1.24, 11.7, 0.42,
                    [{"text": subtitle, "size": 16, "color": MUTED,
                      "italic": True, "space_after": 0}])


def background(slide, color=WHITE):
    bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, Emu(int(SW * 914400)), Emu(int(SH * 914400))
    )
    fill(bg, color)
    no_line(bg)
    no_shadow(bg)
    return bg


def add_notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def add_label(slide, l, t, w, text, *, color=MUTED, size=13, bold=True,
              align=PP_ALIGN.LEFT):
    add_textbox(slide, l, t, w, 0.3,
                [{"text": text, "size": size, "bold": bold, "color": color,
                  "align": align, "space_after": 0}])


# ---------------------------------------------------------------------------
# Slide 1 - Title
# ---------------------------------------------------------------------------
def slide_title(prs):
    slide = new_slide(prs)
    background(slide, NAVY)
    add_rect(slide, 0, 0, 0.22, SH, TEAL)

    add_textbox(slide, 1.1, 1.55, 11.0, 2.1,
                [{"text": "AI Clinical Report", "size": 46, "bold": True,
                  "color": WHITE, "space_after": 0},
                 {"text": "Summarization Assistant", "size": 46, "bold": True,
                  "color": TEAL, "space_after": 0}])

    add_rect(slide, 1.14, 4.32, 1.9, 0.07, TEAL)
    add_textbox(slide, 1.1, 4.52, 6.0, 0.55,
                [{"text": "Team 1", "size": 25, "bold": True, "color": WHITE,
                  "space_after": 0}])
    add_notes(
        slide,
        "Title: AI Clinical Report Summarization Assistant. Team 1. "
        "The deck moves from the business problem to components, five technical "
        "sections, and a short validation statement.",
    )
    return slide


# ---------------------------------------------------------------------------
# Slide 2 - Business problem and solution
# ---------------------------------------------------------------------------
def slide_problem(prs):
    slide = new_slide(prs)
    background(slide, WHITE)
    add_header(slide, "The business problem",
               "Report review is slow, inconsistent, and difficult to audit.")

    add_box(slide, 0.6, 1.82, 5.25, 4.25, [], color=LIGHT, line=BORDER,
            radius=0.06)
    add_textbox(slide, 0.88, 1.98, 4.7, 0.45,
                [{"text": "Today", "size": 20, "bold": True, "color": NAVY,
                  "space_after": 0}])

    problems = [
        "Read every report page by page.",
        "Check units and ranges by eye.",
        "Reconstruct the source chain later.",
    ]
    y = 2.65
    for number, problem in enumerate(problems, start=1):
        add_badge(slide, 0.9, y + 0.02, 0.42, number, color=NAVY_SOFT,
                  fsize=14)
        add_textbox(slide, 1.5, y - 0.03, 4.05, 0.85,
                    [{"text": problem, "size": 18, "color": TEXT,
                      "space_after": 0}])
        y += 1.0

    add_arrow(slide, 5.95, 3.95, 6.25, 3.95)

    add_box(slide, 6.35, 1.82, 6.35, 4.25, [], color=TEAL_LIGHT, line=TEAL,
            radius=0.06)
    add_textbox(slide, 6.63, 1.98, 5.7, 0.45,
                [{"text": "Our solution", "size": 20, "bold": True,
                  "color": NAVY, "space_after": 0}])

    solutions = [
        ("Validate source", "page + exact excerpt"),
        ("Retrieve + summarise", "local guidance, deterministic facts"),
        ("Recommend + review", "cite or abstain; a human decides"),
    ]
    y = 2.55
    for number, (head, sub) in enumerate(solutions, start=1):
        add_badge(slide, 6.65, y + 0.02, 0.4, number, color=TEAL, fsize=13.5)
        add_textbox(slide, 7.22, y - 0.05, 5.1, 0.85,
                    [{"text": head, "size": 18, "bold": True, "color": NAVY,
                      "space_after": 1},
                     {"text": sub, "size": 15, "color": MUTED,
                      "space_after": 0}])
        y += 1.0

    add_box(slide, 0.6, 6.3, 12.1, 0.68,
            [{"text": "A traceable draft, not an autonomous diagnosis.",
              "size": 19, "bold": True, "color": NAVY,
              "align": PP_ALIGN.CENTER, "space_after": 0}],
            color=LIGHT, line=TEAL, radius=0.08)
    add_notes(
        slide,
        "The business problem is a review process that is manual, variable, and "
        "hard to reconstruct. The solution is a source-linked workflow: validate "
        "the report, retrieve local guidance, assemble only validated facts, and "
        "stop for an explicit human decision. The next slides show the components "
        "and the five technical sections.",
    )
    add_footer(slide, 2)
    return slide


# ---------------------------------------------------------------------------
# Slide 3 - Architecture
# ---------------------------------------------------------------------------
def slide_architecture(prs):
    slide = new_slide(prs)
    background(slide, WHITE)
    add_header(slide, "Architecture",
               "Components, boundaries, and three bounded agents.")

    add_box(slide, 0.6, 1.72, 12.1, 0.75, [], color=NAVY, line=NAVY,
            radius=0.06)
    add_textbox(slide, 0.86, 1.88, 4.0, 0.42,
                [{"text": "Streamlit UI", "size": 20, "bold": True,
                  "color": WHITE, "space_after": 0}])
    add_textbox(slide, 4.8, 1.95, 7.65, 0.34,
                [{"text": "upload  \u00b7  run  \u00b7  evidence  \u00b7  review",
                  "size": 14.5, "color": PALE_TEXT, "align": PP_ALIGN.RIGHT,
                  "space_after": 0}])

    add_arrow(slide, 6.65, 2.47, 6.65, 2.67)
    add_box(slide, 0.6, 2.7, 12.1, 0.75, [], color=NAVY_SOFT,
            line=NAVY_SOFT, radius=0.06)
    add_textbox(slide, 0.86, 2.86, 4.0, 0.42,
                [{"text": "CrewAI Flow", "size": 20, "bold": True,
                  "color": WHITE, "space_after": 0}])
    add_textbox(slide, 4.6, 2.93, 7.85, 0.34,
                [{"text": "resolve  \u00b7  fork  \u00b7  join  \u00b7  route",
                  "size": 14.5, "color": PALE_TEXT, "align": PP_ALIGN.RIGHT,
                  "space_after": 0}])

    centers = [1.98, 4.88, 7.78, 10.98]
    for cx in centers:
        add_arrow(slide, cx, 3.45, cx, 3.72)

    agents = [
        (0.6, 2.75, "Report Analysis", "extract \u00b7 validate \u00b7 compare",
         "Gemini call", TEAL_LIGHT, TEAL, AMBER_DARK),
        (3.5, 2.75, "Summary", "IDs in \u00b7 facts out",
         "Gemini call", TEAL_LIGHT, TEAL, AMBER_DARK),
        (6.4, 2.75, "Recommendation", "cite both sides \u00b7 or abstain",
         "Gemini call", TEAL_LIGHT, TEAL, AMBER_DARK),
        (9.3, 3.4, "Guideline RAG", "chunks \u00b7 embeddings \u00b7 Chroma",
         "local \u00b7 no API", LIGHT, RGBColor(0x9F, 0xB3, 0xC6), TEAL_DARK),
    ]
    for x, w, head, sub, tag, box_fill, box_line, tag_color in agents:
        add_box(slide, x, 3.75, w, 1.45,
                [{"text": head, "size": 18, "bold": True, "color": NAVY,
                  "space_after": 2},
                 {"text": sub, "size": 14, "color": MUTED, "space_after": 7},
                 {"text": tag, "size": 13.5, "bold": True, "color": tag_color,
                  "space_after": 0}],
                color=box_fill, line=box_line, radius=0.08)

    support = [
        (0.6, 3.5, "PDF ingestion", "in-memory gates"),
        (4.4, 3.7, "Pydantic boundaries", "schema + evidence"),
        (8.4, 4.3, "SQLite store", "immutable + review"),
    ]
    for x, w, head, sub in support:
        add_box(slide, x, 5.62, w, 0.82,
                [{"text": head, "size": 17, "bold": True, "color": NAVY,
                  "space_after": 1},
                 {"text": sub, "size": 14, "color": MUTED, "space_after": 0}],
                color=LIGHT, line=BORDER, radius=0.08)

    add_arrow(slide, 2.35, 5.62, 2.35, 5.22, color=TEAL_DARK, width=1.8)
    add_label(slide, 2.48, 5.3, 1.2, "pages", color=TEAL_DARK, size=14,
              bold=False)

    add_textbox(slide, 0.6, 6.62, 12.1, 0.3,
                [{"text": "Teal = agent   \u00b7   amber = Gemini boundary   "
                          "\u00b7   grey = deterministic support",
                  "size": 13.5, "italic": True, "color": MUTED,
                  "space_after": 0}])
    add_notes(
        slide,
        "The UI is a thin action layer. CrewAI Flow handles sequencing, the fork, "
        "the join, and routing. The three teal boxes are the only AI agents. "
        "Guideline RAG, PDF ingestion, Pydantic validation, SQLite, and session "
        "state are supporting components. The orchestration slide shows the "
        "execution order.",
    )
    add_footer(slide, 3)
    return slide


# ---------------------------------------------------------------------------
# Slide 4 - Ingestion + analysis
# ---------------------------------------------------------------------------
def slide_ingest_analysis(prs):
    slide = new_slide(prs)
    background(slide, WHITE)
    add_header(slide, "Ingestion + analysis",
               "From PDF bytes to source-verified findings.")

    add_box(slide, 0.6, 1.76, 4.3, 5.0, [], color=LIGHT, line=BORDER,
            radius=0.06)
    add_textbox(slide, 0.86, 1.92, 3.7, 0.4,
                [{"text": "PDF intake", "size": 20, "bold": True,
                  "color": NAVY, "space_after": 0}])

    gates = [
        ("Size & type", "\u2264 10 MiB  \u00b7  %PDF"),
        ("Parse", "pypdf  \u00b7  page text"),
        ("Identity", "SHA-256  \u00b7  in memory"),
        ("No text", "OCR warning only"),
    ]
    y = 2.55
    for head, sub in gates:
        add_rect(slide, 0.86, y + 0.02, 0.07, 0.65, TEAL)
        add_textbox(slide, 1.08, y - 0.02, 3.5, 0.8,
                    [{"text": head, "size": 18, "bold": True, "color": NAVY,
                      "space_after": 1},
                     {"text": sub, "size": 14.5, "color": MUTED,
                      "space_after": 0}])
        y += 0.95

    add_textbox(slide, 0.86, 6.35, 3.7, 0.3,
                [{"text": "No PDF bytes written to disk.", "size": 15,
                  "bold": True, "color": TEAL_DARK, "space_after": 0}])

    steps = [
        ("Extract", "findings"),
        ("Validate", "page + excerpt"),
        ("Compare", "Decimal"),
    ]
    x = 5.25
    for i, (head, sub) in enumerate(steps):
        add_box(slide, x, 1.85, 2.25, 0.9,
                [{"text": head, "size": 18, "bold": True, "color": NAVY,
                  "space_after": 1},
                 {"text": sub, "size": 14.5, "color": MUTED,
                  "space_after": 0}],
                color=CARD, line=TEAL if i == 2 else BORDER, radius=0.08)
        if i < 2:
            add_arrow(slide, x + 2.25, 2.3, x + 2.60, 2.3)
        x += 2.60

    add_box(slide, 5.25, 3.05, 7.45, 1.55, [], color=CARD, line=BORDER,
            radius=0.06)
    add_textbox(slide, 5.48, 3.14, 4.0, 0.3,
                [{"text": "Verified synthetic example", "size": 14,
                  "bold": True, "color": MUTED, "space_after": 0}])
    add_textbox(slide, 5.48, 3.46, 2.55, 0.6,
                [{"text": "110 mg/dL", "size": 30, "bold": True,
                  "color": NAVY, "space_after": 0}])
    add_textbox(slide, 8.08, 3.56, 2.0, 0.4,
                [{"text": "vs 70\u201399 mg/dL", "size": 18, "color": MUTED,
                  "space_after": 0}])
    add_pill(slide, 10.3, 3.5, 2.1, 0.45, "above_range",
             color=TEAL_LIGHT, line=TEAL, text_color=TEAL_DARK, size=16)

    add_label(slide, 5.25, 4.78, 3.4, "SEPARATED OUTPUTS", color=MUTED, size=14)
    outputs = ["Patient", "Narrative", "Control", "Warnings"]
    x = 5.25
    for out in outputs:
        add_pill(slide, x, 5.08, 1.7, 0.5, out, color=LIGHT, line=BORDER,
                 text_color=NAVY, size=15)
        x += 1.85

    add_box(slide, 5.25, 5.78, 7.45, 0.65,
            [{"text": "The model proposes.  Code checks.", "size": 20,
              "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER,
              "space_after": 0}],
            color=NAVY, line=NAVY, radius=0.08)

    add_textbox(slide, 5.25, 6.55, 7.45, 0.32,
                [{"text": "Missing range \u2192 not assessed   \u00b7   "
                          "unit mismatch \u2192 review",
                  "size": 15, "italic": True, "color": MUTED,
                  "space_after": 0}])
    add_notes(
        slide,
        "This section covers the input boundary and the first agent. The PDF is "
        "validated before parsing; text is extracted page by page and hashed. The "
        "Report Analysis Agent proposes findings, but page evidence, units, and "
        "the numeric comparison are checked in code. The verified synthetic "
        "glucose example is 110 mg/dL against a supplied 70-99 mg/dL interval, "
        "which is above_range. A missing range is not inferred. "
        "Suggested section time: 2 minutes.",
    )
    add_footer(slide, 4)
    return slide


# ---------------------------------------------------------------------------
# Slide 5 - Guideline RAG
# ---------------------------------------------------------------------------
def slide_guideline_rag(prs):
    slide = new_slide(prs)
    background(slide, WHITE)
    add_header(slide, "Guideline RAG",
               "Local retrieval with exact page provenance.")

    pipeline = [
        ("Guideline PDF", "selectable text"),
        ("Chunks", "800 / 120"),
        ("Embeddings", "384-d local"),
        ("Chroma", "persistent"),
        ("Passages", "page + hash"),
    ]
    x = 0.6
    for i, (head, sub) in enumerate(pipeline):
        add_box(slide, x, 1.8, 2.15, 1.2,
                [{"text": head, "size": 18, "bold": True, "color": NAVY,
                  "space_after": 3},
                 {"text": sub, "size": 14.5, "color": MUTED,
                  "space_after": 0}],
                color=TEAL_LIGHT if i in (1, 2) else LIGHT,
                line=TEAL if i in (1, 2) else BORDER, radius=0.08)
        if i < 4:
            add_arrow(slide, x + 2.15, 2.4, x + 2.45, 2.4)
        x += 2.45

    add_box(slide, 0.6, 3.35, 7.2, 1.3, [], color=CARD, line=BORDER,
            radius=0.06)
    add_textbox(slide, 0.84, 3.48, 4.0, 0.35,
                [{"text": "Every chunk knows", "size": 18, "bold": True,
                  "color": NAVY, "space_after": 0}])
    tags = ["document", "page", "hash"]
    x = 0.84
    for tag in tags:
        add_pill(slide, x, 3.95, 1.9, 0.4, tag, color=LIGHT, line=BORDER,
                 text_color=MUTED, size=15, bold=False)
        x += 2.05

    add_box(slide, 8.1, 3.35, 4.6, 1.3, [], color=PLUM_LIGHT, line=PLUM,
            radius=0.06)
    add_textbox(slide, 8.34, 3.46, 3.0, 0.35,
                [{"text": "Index fingerprint", "size": 18, "bold": True,
                  "color": NAVY, "space_after": 0}])
    add_textbox(slide, 8.34, 3.84, 2.4, 0.43,
                [{"text": "SHA-256", "size": 23, "bold": True, "color": PLUM,
                  "space_after": 0}])
    add_textbox(slide, 10.45, 3.92, 2.05, 0.55,
                [{"text": "config + document hashes", "size": 14,
                  "color": MUTED, "space_after": 0}])

    add_box(slide, 0.6, 4.95, 7.2, 1.45, [], color=LIGHT, line=BORDER,
            radius=0.06)
    add_label(slide, 0.84, 5.08, 5.2, "EXACT SYNTHETIC PASSAGE  \u00b7  PAGE 2",
              color=TEAL_DARK, size=14)
    add_textbox(slide, 0.84, 5.38, 6.7, 0.45,
                [{"text": "\u201cDo not infer an interval.\u201d", "size": 22,
                  "italic": True, "color": NAVY, "space_after": 0}])
    add_textbox(slide, 0.84, 5.95, 6.5, 0.3,
                [{"text": "Missing reference-range review", "size": 14.5,
                  "color": MUTED, "space_after": 0}])

    add_box(slide, 8.1, 4.95, 4.6, 1.45,
            [{"text": "Similarity ranks relevance.", "size": 18, "bold": True,
              "color": NAVY, "space_after": 3},
             {"text": "It is not confidence.", "size": 16,
              "color": MUTED, "space_after": 0}],
            color=AMBER_LIGHT, line=AMBER, radius=0.08)

    add_textbox(slide, 0.6, 6.62, 12.1, 0.3,
                [{"text": "Local embeddings  \u00b7  persistent Chroma  \u00b7  "
                          "no Gemini call in retrieval",
                  "size": 14.5, "italic": True, "color": MUTED,
                  "align": PP_ALIGN.CENTER, "space_after": 0}])
    add_notes(
        slide,
        "This section covers the guideline path. Pages become overlapping chunks "
        "with deterministic IDs; local embeddings go into persistent Chroma. "
        "Retrieved passages keep document, page, and hash provenance. The quoted "
        "line is an exact sentence from the synthetic guideline on page 2. "
        "Similarity is a ranking signal, not clinical confidence. "
        "Suggested section time: 2 minutes.",
    )
    add_footer(slide, 5)
    return slide


# ---------------------------------------------------------------------------
# Slide 6 - Summary
# ---------------------------------------------------------------------------
def slide_summary(prs):
    slide = new_slide(prs)
    background(slide, WHITE)
    add_header(slide, "Summary",
               "Gemini selects IDs; deterministic code writes the facts.")

    add_box(slide, 0.6, 1.85, 3.6, 3.2, [], color=AMBER_LIGHT, line=AMBER,
            radius=0.06)
    add_textbox(slide, 0.86, 2.0, 3.0, 0.4,
                [{"text": "Gemini", "size": 20, "bold": True, "color": NAVY,
                  "space_after": 0}])
    add_textbox(slide, 0.86, 2.38, 3.0, 0.3,
                [{"text": "bounded ID selection", "size": 15,
                  "color": MUTED, "space_after": 0}])
    ids = ["overview", "important", "repeats", "narrative", "uncertainty"]
    y = 2.82
    for tag in ids:
        add_pill(slide, 0.86, y, 3.05, 0.34, tag, color=CARD, line=AMBER,
                 text_color=NAVY, size=14.5, bold=False)
        y += 0.4

    add_arrow(slide, 4.2, 3.45, 4.55, 3.45)

    add_box(slide, 4.6, 1.85, 3.3, 3.2, [], color=NAVY, line=NAVY,
            radius=0.06)
    add_textbox(slide, 4.86, 2.0, 2.9, 0.4,
                [{"text": "Assembler", "size": 20, "bold": True,
                  "color": WHITE, "space_after": 0}])
    add_textbox(slide, 4.86, 2.38, 2.9, 0.3,
                [{"text": "deterministic", "size": 15, "color": TEAL,
                  "space_after": 0}])
    lines = ["values \u00b7 units \u00b7 ranges", "repeat groups",
             "pages \u00b7 excerpts", "uncertainties"]
    y = 2.9
    for line in lines:
        add_textbox(slide, 4.86, y, 2.85, 0.32,
                    [{"text": BULLET + line, "size": 16, "color": PALE_TEXT,
                      "space_after": 0}])
        y += 0.48

    add_arrow(slide, 7.9, 3.45, 8.25, 3.45)

    add_box(slide, 8.3, 1.85, 4.4, 3.2, [], color=CARD, line=TEAL,
            radius=0.06)
    add_textbox(slide, 8.56, 2.0, 3.8, 0.4,
                [{"text": "Summary", "size": 20, "bold": True,
                  "color": NAVY, "space_after": 0}])
    sections = ["Overview", "Important findings", "Repeat measurements",
                "Narrative context", "Uncertainties"]
    y = 2.62
    for section in sections:
        add_box(slide, 8.56, y, 3.85, 0.36, [],
                color=TEAL_LIGHT, line=TEAL_LIGHT, radius=0.5)
        add_textbox(slide, 8.72, y + 0.04, 3.55, 0.3,
                    [{"text": section, "size": 14.5, "bold": True,
                      "color": NAVY, "space_after": 0}])
        y += 0.44

    add_box(slide, 0.6, 5.3, 6.2, 0.8,
            [{"text": "IDs in.  Facts out.", "size": 24, "bold": True,
              "color": NAVY, "space_after": 0}],
            color=LIGHT, line=TEAL, radius=0.08)
    add_textbox(slide, 7.0, 5.39, 5.7, 0.6,
                [{"text": "The model never writes the numeric fact.",
                  "size": 16, "color": MUTED, "space_after": 0}])

    add_textbox(slide, 0.6, 6.4, 12.1, 0.32,
                [{"text": "Repeat groups: matching analyte + unit; no invented "
                          "chronology.",
                  "size": 15, "italic": True, "color": MUTED,
                  "align": PP_ALIGN.CENTER, "space_after": 0}])
    add_notes(
        slide,
        "This section covers the summary. Gemini returns finding IDs only. "
        "Deterministic code rebuilds values, repeat groups, evidence, and "
        "mandatory uncertainties from the validated analysis. The model never "
        "writes the numeric fact it claims to summarise. "
        "Suggested section time: 2 minutes.",
    )
    add_footer(slide, 6)
    return slide


# ---------------------------------------------------------------------------
# Slide 7 - Recommendation + human review
# ---------------------------------------------------------------------------
def slide_recommendation(prs):
    slide = new_slide(prs)
    background(slide, WHITE)
    add_header(slide, "Recommendation + human review",
               "Every suggestion cites both sides \u2014 or declines.")

    add_box(slide, 0.6, 1.8, 3.1, 2.8, [], color=LIGHT, line=BORDER,
            radius=0.06)
    add_textbox(slide, 0.84, 1.94, 2.6, 0.4,
                [{"text": "Inputs", "size": 20, "bold": True, "color": NAVY,
                  "space_after": 0}])
    inputs = ["patient findings", "validated summary", "guideline passages"]
    y = 2.4
    for item in inputs:
        add_pill(slide, 0.84, y, 2.6, 0.48, item, color=CARD, line=BORDER,
                 text_color=NAVY, size=15, bold=False)
        y += 0.64

    add_arrow(slide, 3.7, 3.2, 4.05, 3.2)

    add_box(slide, 4.1, 2.0, 2.8, 2.4,
            [{"text": "Recommendation", "size": 20, "bold": True,
              "color": WHITE, "space_after": 0},
             {"text": "Agent", "size": 20, "bold": True, "color": WHITE,
              "space_after": 6},
             {"text": "bounded Gemini call", "size": 14.5, "color": AMBER,
              "space_after": 0}],
            color=NAVY, line=NAVY, radius=0.08)

    add_arrow(slide, 6.9, 2.75, 7.35, 2.4)
    add_arrow(slide, 6.9, 3.55, 7.95, 3.7)

    add_box(slide, 7.4, 1.8, 5.3, 1.25,
            [{"text": "Suggestion", "size": 20, "bold": True, "color": NAVY,
              "space_after": 1},
             {"text": "finding IDs + passage IDs", "size": 16,
              "color": MUTED, "space_after": 4},
             {"text": "patient evidence + guideline evidence", "size": 15,
              "bold": True, "color": TEAL_DARK, "space_after": 0}],
            color=TEAL_LIGHT, line=TEAL, radius=0.08)

    add_box(slide, 8.0, 3.15, 4.3, 1.15,
            [{"text": "Abstain", "size": 20, "bold": True, "color": NAVY,
              "space_after": 1},
             {"text": "absent \u00b7 unrelated \u00b7 contradictory \u00b7 "
                      "insufficient", "size": 14, "color": MUTED,
              "space_after": 0}],
            color=AMBER_LIGHT, line=AMBER, radius=0.08)

    add_arrow(slide, 12.4, 3.05, 12.4, 4.62, color=TEAL_DARK, width=1.8)

    add_box(slide, 0.6, 4.7, 12.1, 1.0, [], color=NAVY, line=NAVY,
            radius=0.06)
    add_textbox(slide, 0.86, 4.9, 3.8, 0.4,
                [{"text": "Human review", "size": 20, "bold": True,
                  "color": WHITE, "space_after": 0}])
    add_textbox(slide, 4.4, 4.95, 8.0, 0.35,
                [{"text": "approve / reject exact draft  \u00b7  edit = new "
                          "pending version", "size": 15, "color": PALE_TEXT,
                  "align": PP_ALIGN.RIGHT, "space_after": 0}])

    add_box(slide, 0.6, 5.95, 12.1, 0.55,
            [{"text": "No diagnosis  \u00b7  no medication changes  \u00b7  "
                      "no doses  \u00b7  no urgency", "size": 16,
              "bold": True, "color": NAVY, "align": PP_ALIGN.CENTER,
              "space_after": 0}],
            color=LIGHT, line=BORDER, radius=0.08)

    add_textbox(slide, 0.6, 6.68, 12.1, 0.3,
                [{"text": "Approval is a review decision, not clinical "
                          "certification.",
                  "size": 15, "italic": True, "color": MUTED,
                  "align": PP_ALIGN.CENTER, "space_after": 0}])
    add_notes(
        slide,
        "This section covers the recommendation and the human decision. Every "
        "suggestion needs both a patient finding and a guideline passage with "
        "exact evidence. If guidance is absent, unrelated, contradictory, or too "
        "thin, the agent abstains. The draft stays pending until a person "
        "approves or rejects that exact version. "
        "Suggested section time: 2 minutes.",
    )
    add_footer(slide, 7)
    return slide


# ---------------------------------------------------------------------------
# Slide 8 - Orchestration + persistence
# ---------------------------------------------------------------------------
def slide_orchestration(prs):
    slide = new_slide(prs)
    background(slide, WHITE)
    add_header(slide, "Orchestration + persistence",
               "Execution order, parallel branches, and exact reuse.")

    add_box(slide, 0.6, 1.78, 5.9, 0.62,
            [{"text": "Analysis", "size": 18, "bold": True, "color": WHITE,
              "space_after": 0}],
            color=NAVY_SOFT, line=NAVY_SOFT, radius=0.08)
    add_arrow(slide, 3.55, 2.4, 3.55, 2.55, width=1.8)
    add_line(slide, 1.95, 2.55, 5.15, 2.55)
    add_arrow(slide, 1.95, 2.55, 1.95, 2.69, width=1.8)
    add_arrow(slide, 5.15, 2.55, 5.15, 2.69, width=1.8)

    add_box(slide, 0.6, 2.7, 2.7, 0.76,
            [{"text": "Summary", "size": 17, "bold": True, "color": NAVY,
              "space_after": 0}],
            color=TEAL_LIGHT, line=TEAL, radius=0.08)
    add_box(slide, 3.8, 2.7, 2.7, 0.76,
            [{"text": "Retrieval", "size": 17, "bold": True, "color": NAVY,
              "space_after": 0}],
            color=LIGHT, line=RGBColor(0x9F, 0xB3, 0xC6), radius=0.08)

    add_arrow(slide, 1.95, 3.46, 1.95, 3.65, width=1.8)
    add_arrow(slide, 5.15, 3.46, 5.15, 3.65, width=1.8)
    add_line(slide, 1.95, 3.65, 5.15, 3.65)
    add_arrow(slide, 3.55, 3.65, 3.55, 3.84, width=1.8)

    add_box(slide, 0.6, 3.85, 5.9, 0.62,
            [{"text": "Join + fingerprint re-check", "size": 18, "bold": True,
              "color": WHITE, "space_after": 0}],
            color=NAVY_SOFT, line=NAVY_SOFT, radius=0.08)
    add_arrow(slide, 3.55, 4.47, 3.55, 4.64)
    add_box(slide, 0.6, 4.65, 5.9, 0.62,
            [{"text": "Recommendation", "size": 18, "bold": True,
              "color": NAVY, "space_after": 0}],
            color=TEAL_LIGHT, line=TEAL, radius=0.08)
    add_arrow(slide, 3.55, 5.27, 3.55, 5.44)
    add_box(slide, 0.6, 5.45, 5.9, 0.62,
            [{"text": "Human review", "size": 18, "bold": True,
              "color": NAVY, "space_after": 0}],
            color=LIGHT, line=BORDER, radius=0.08)

    add_box(slide, 7.2, 1.78, 5.5, 0.9,
            [{"text": "SQLite", "size": 20, "bold": True, "color": WHITE,
              "space_after": 1},
             {"text": "immutable local results", "size": 14,
              "color": PALE_TEXT, "space_after": 0}],
            color=TEAL_DARK, line=TEAL_DARK, radius=0.06)
    tags = ["immutable records", "exact compatibility keys", "revalidate on load"]
    y = 2.92
    for tag in tags:
        add_box(slide, 7.2, y, 5.5, 0.5,
                [{"text": tag, "size": 15, "color": NAVY, "space_after": 0}],
                color=LIGHT, line=BORDER, radius=0.08)
        y += 0.6

    add_box(slide, 7.2, 4.85, 5.5, 0.85,
            [{"text": "No new Gemini call for a reused stage.", "size": 18,
              "bold": True, "color": NAVY, "space_after": 2},
             {"text": "Exact keys + validated load.", "size": 14.5,
              "color": MUTED, "space_after": 0}],
            color=TEAL_LIGHT, line=TEAL, radius=0.06)

    routes = ["reuse", "generate", "retrieve", "abstain", "await review",
              "approved", "rejected"]
    widths = [1.2, 1.4, 1.4, 1.4, 1.8, 1.5, 1.5]
    route_total = sum(widths) + 0.1 * (len(widths) - 1)
    x = (SW - route_total) / 2
    for label, w in zip(routes, widths):
        add_pill(slide, x, 6.15, w, 0.38, label, color=CARD, line=BORDER,
                 text_color=MUTED, size=14, bold=False)
        x += w + 0.1

    add_textbox(slide, 0.6, 6.68, 12.1, 0.3,
                [{"text": "A failed branch preserves every saved upstream stage.",
                  "size": 14.5, "italic": True, "color": MUTED,
                  "align": PP_ALIGN.CENTER, "space_after": 0}])
    add_notes(
        slide,
        "This section covers execution order and routing. Analysis resolves "
        "first. Summary and local retrieval then run as two parallel branches. "
        "Both branches join, and the guideline fingerprint is rechecked before "
        "Recommendation. The route is recorded as reuse, generate, retrieve, "
        "abstain, await review, approved, or rejected. A compatible reused stage "
        "does not make a new Gemini call. A failed branch preserves any stage "
        "that was already saved. Suggested section time: 2 minutes.",
    )
    add_footer(slide, 8)
    return slide


# ---------------------------------------------------------------------------
# Slide 9 - Validation and limitations
# ---------------------------------------------------------------------------
def slide_validation(prs):
    slide = new_slide(prs)
    background(slide, WHITE)
    add_header(slide, "Validation and limitations",
               "Offline evidence is explicit; live behavior is checked separately.")

    validation = [
        ("173", "tests passed", "1 optional test skipped; Ruff + lockfile clean"),
        ("0", "API calls in mocked demo", "No network, key, or model download"),
        ("1", "manual live smoke request", "Passed separately from offline tests"),
    ]
    x = 0.6
    for number, head, sub in validation:
        add_box(slide, x, 1.85, 3.9, 1.95, [], color=LIGHT, line=BORDER,
                radius=0.08)
        add_textbox(slide, x + 0.22, 2.0, 1.2, 0.7,
                    [{"text": number, "size": 38, "bold": True, "color": TEAL,
                      "space_after": 0}])
        add_textbox(slide, x + 1.45, 2.08, 2.25, 0.7,
                    [{"text": head, "size": 18, "bold": True, "color": NAVY,
                      "space_after": 0}])
        add_textbox(slide, x + 0.22, 2.85, 3.45, 0.75,
                    [{"text": sub, "size": 14.5, "color": MUTED,
                      "space_after": 0}])
        x += 4.1

    add_box(slide, 0.6, 4.1, 6.0, 1.85, [], color=AMBER_LIGHT, line=AMBER,
            radius=0.06)
    add_textbox(slide, 0.86, 4.24, 5.5, 0.4,
                [{"text": "Limitation: not a medical device", "size": 20,
                  "bold": True, "color": NAVY, "space_after": 0}])
    add_textbox(slide, 0.86, 4.82, 5.45, 0.85,
                [{"text": "Synthetic educational demo only. No diagnosis, "
                          "treatment, or clinical advice.", "size": 17,
                  "color": TEXT, "space_after": 0}])

    add_box(slide, 6.7, 4.1, 6.0, 1.85, [], color=AMBER_LIGHT, line=AMBER,
            radius=0.06)
    add_textbox(slide, 6.96, 4.24, 5.5, 0.4,
                [{"text": "Limitation: PDFs and similarity", "size": 20,
                  "bold": True, "color": NAVY, "space_after": 0}])
    add_textbox(slide, 6.96, 4.82, 5.45, 0.85,
                [{"text": "Scanned PDFs need OCR, which is out of scope. "
                          "Similarity is not confidence.", "size": 17,
                  "color": TEXT, "space_after": 0}])

    add_box(slide, 0.6, 6.2, 12.1, 0.68,
            [{"text": "Show the source.  Decline when unsure.  Let a human decide.",
              "size": 18, "bold": True, "color": WHITE,
              "align": PP_ALIGN.CENTER, "space_after": 0}],
            color=NAVY, line=NAVY, radius=0.08)
    add_notes(
        slide,
        "Validation has three separate statements. The offline suite has 173 "
        "passing tests and one optional real-embedding test skipped by default; "
        "Ruff and the lockfile checks also pass. The mocked offline demo makes "
        "zero API calls and uses no network, key, or model download. A separate "
        "manual live smoke check made one synthetic Gemini request and passed; it "
        "is not part of the offline suite. Two limitations remain: this is a "
        "synthetic educational demo, not a medical device; and scanned PDFs need "
        "OCR, which is out of scope, while similarity is not confidence.",
    )
    add_footer(slide, 9)
    return slide


# ---------------------------------------------------------------------------
SLIDE_BUILDERS = [
    slide_title,
    slide_problem,
    slide_architecture,
    slide_ingest_analysis,
    slide_guideline_rag,
    slide_summary,
    slide_recommendation,
    slide_orchestration,
    slide_validation,
]


def build(out_path: Path) -> Path:
    global TOTAL
    TOTAL = len(SLIDE_BUILDERS)

    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)

    for builder in SLIDE_BUILDERS:
        builder(prs)

    prs.core_properties.title = "AI Clinical Report Summarization Assistant"
    prs.core_properties.author = "Team 1"
    prs.core_properties.comments = (
        "9-slide visual overview: business problem, architecture, five technical "
        "sections, validation and limitations."
    )

    prs.save(str(out_path))
    return out_path


if __name__ == "__main__":
    root = Path(__file__).resolve().parent.parent
    target = root / "Clinical_Report_Assistant_Presentation.pptx"
    build(target)
    print(f"[OK] wrote {target}")

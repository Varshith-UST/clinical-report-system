"""Build the 12-slide capstone presentation as an editable PowerPoint file.

Run with:
    uv run --with python-pptx python scripts/build_capstone_deck.py
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

# ---------------------------------------------------------------------------
# Theme
# ---------------------------------------------------------------------------

W, H = 13.333, 7.5
MX = 0.55
CW = 12.233

BG = RGBColor(0x0B, 0x12, 0x20)
CARD = RGBColor(0x12, 0x1C, 0x2E)
CARD_ALT = RGBColor(0x0F, 0x18, 0x28)
BORDER = RGBColor(0x25, 0x34, 0x49)
TEXT = RGBColor(0xE8, 0xEE, 0xF7)
MUTED = RGBColor(0x96, 0xA7, 0xBE)
DIM = RGBColor(0x6B, 0x7C, 0x93)
TEAL = RGBColor(0x2D, 0xD4, 0xBF)
BLUE = RGBColor(0x60, 0xA5, 0xFA)
AMBER = RGBColor(0xFB, 0xBF, 0x24)
PURPLE = RGBColor(0xA7, 0x8B, 0xFA)
GREEN = RGBColor(0x34, 0xD3, 0x99)
RED = RGBColor(0xF8, 0x71, 0x71)

FONT = "Segoe UI"


# ---------------------------------------------------------------------------
# Text helpers
# ---------------------------------------------------------------------------


def p(
    text: str,
    *,
    sz: float = 11,
    b: bool = False,
    c: RGBColor = TEXT,
    i: bool = False,
    align: PP_ALIGN | None = None,
    after: float = 2,
    before: float = 0,
    line: float = 1.0,
    font: str = FONT,
) -> dict:
    return {
        "runs": [(text, {"sz": sz, "b": b, "c": c, "i": i, "font": font})],
        "align": align,
        "after": after,
        "before": before,
        "line": line,
    }


def rp(
    runs: list[tuple[str, dict]],
    *,
    align: PP_ALIGN | None = None,
    after: float = 2,
    before: float = 0,
    line: float = 1.0,
) -> dict:
    return {
        "runs": runs,
        "align": align,
        "after": after,
        "before": before,
        "line": line,
    }


def bullet(
    text: str,
    *,
    sz: float = 10.5,
    c: RGBColor = TEXT,
    bc: RGBColor = TEAL,
    b: bool = False,
    after: float = 5,
    glyph: str = "\u2022",
) -> dict:
    return {
        "runs": [
            (glyph + "  ", {"sz": sz, "b": True, "c": bc}),
            (text, {"sz": sz, "b": b, "c": c}),
        ],
        "after": after,
    }


def write_tf(
    tf,
    paras: list[dict],
    *,
    anchor=MSO_ANCHOR.TOP,
    margins: tuple[float, float, float, float] = (0, 0, 0, 0),
) -> None:
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Inches(margins[0])
    tf.margin_top = Inches(margins[1])
    tf.margin_right = Inches(margins[2])
    tf.margin_bottom = Inches(margins[3])
    first = True
    for spec in paras:
        para = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        if spec.get("align") is not None:
            para.alignment = spec["align"]
        para.space_before = Pt(spec.get("before", 0))
        para.space_after = Pt(spec.get("after", 2))
        para.line_spacing = spec.get("line", 1.0)
        for text, style in spec["runs"]:
            run = para.add_run()
            run.text = text
            f = run.font
            f.name = style.get("font", FONT)
            f.size = Pt(style.get("sz", 11))
            f.bold = style.get("b", False)
            f.italic = style.get("i", False)
            f.color.rgb = style.get("c", TEXT)


def add_text(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    paras: list[dict],
    *,
    anchor=MSO_ANCHOR.TOP,
    margins: tuple[float, float, float, float] = (0, 0, 0, 0),
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    write_tf(box.text_frame, paras, anchor=anchor, margins=margins)
    return box


# ---------------------------------------------------------------------------
# Shape helpers
# ---------------------------------------------------------------------------


def rect(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    fill: RGBColor | None = CARD,
    line: RGBColor | None = BORDER,
    shape=MSO_SHAPE.ROUNDED_RECTANGLE,
    radius: float = 0.08,
    line_w: float = 1.0,
    dash: MSO_LINE_DASH_STYLE | None = None,
):
    shp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        shp.adjustments[0] = radius
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
        if dash is not None:
            shp.line.dash_style = dash
    shp.shadow.inherit = False
    return shp


def card_text(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    paras: list[dict],
    *,
    anchor=MSO_ANCHOR.TOP,
    margins: tuple[float, float, float, float] = (0.16, 0.12, 0.16, 0.10),
):
    return add_text(slide, x, y, w, h, paras, anchor=anchor, margins=margins)


def chip(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str,
    *,
    fill: RGBColor = CARD,
    line: RGBColor = BORDER,
    color: RGBColor = TEXT,
    sz: float = 8.5,
    bold: bool = False,
):
    shp = rect(slide, x, y, w, h, fill=fill, line=line, radius=0.5, line_w=0.75)
    write_tf(
        shp.text_frame,
        [p(text, sz=sz, b=bold, c=color, align=PP_ALIGN.CENTER, after=0)],
        anchor=MSO_ANCHOR.MIDDLE,
        margins=(0.04, 0.01, 0.04, 0.01),
    )
    return shp


def badge(slide, x: float, y: float, d: float, text: str, *, color=TEAL):
    shp = rect(
        slide,
        x,
        y,
        d,
        d,
        fill=RGBColor(0x0D, 0x2A, 0x2C),
        line=color,
        shape=MSO_SHAPE.OVAL,
        line_w=0.9,
    )
    write_tf(
        shp.text_frame,
        [p(text, sz=9.5, b=True, c=color, align=PP_ALIGN.CENTER, after=0)],
        anchor=MSO_ANCHOR.MIDDLE,
        margins=(0, 0, 0, 0),
    )
    return shp


def arrowhead(connector, *, size: str = "med"):
    line = connector.line
    ln = line._get_or_add_ln()
    tail = ln.find("{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd")
    if tail is None:
        tail = OxmlElement("a:tailEnd")
        ln.append(tail)
    tail.set("type", "triangle")
    tail.set("w", size)
    tail.set("len", size)


def arrow(
    slide,
    x1: float,
    y1: float,
    x2: float,
    y2: float,
    *,
    color: RGBColor = DIM,
    width: float = 1.5,
    dash: MSO_LINE_DASH_STYLE | None = None,
):
    conn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    conn.line.color.rgb = color
    conn.line.width = Pt(width)
    if dash is not None:
        conn.line.dash_style = dash
    arrowhead(conn)
    return conn


def down_arrow(slide, cx: float, y: float, h: float, *, color=TEAL, w: float = 0.16):
    shp = rect(
        slide,
        cx - w / 2,
        y,
        w,
        h,
        fill=color,
        line=None,
        shape=MSO_SHAPE.DOWN_ARROW,
    )
    return shp


def right_arrow_shape(slide, x, y, w=0.26, h=0.18, *, color=DIM):
    return rect(slide, x, y, w, h, fill=color, line=None, shape=MSO_SHAPE.RIGHT_ARROW)


# ---------------------------------------------------------------------------
# Slide scaffolding
# ---------------------------------------------------------------------------


def base_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG

    top = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.05))
    top.fill.gradient()
    stops = top.fill.gradient_stops
    stops[0].color.rgb = TEAL
    stops[0].position = 0.0
    stops[1].color.rgb = BLUE
    stops[1].position = 1.0
    top.fill.gradient_angle = 0
    top.line.fill.background()
    top.shadow.inherit = False
    return slide


def header(slide, kicker: str, title: str, subtitle: str | None = None):
    rect(slide, MX, 0.34, 0.055, 0.26, fill=TEAL, line=None, shape=MSO_SHAPE.RECTANGLE)
    add_text(
        slide,
        MX + 0.16,
        0.28,
        9.0,
        0.28,
        [p(kicker.upper(), sz=9, b=True, c=TEAL, after=0)],
    )
    add_text(
        slide,
        MX,
        0.55,
        CW,
        0.58,
        [p(title, sz=27, b=True, c=TEXT, after=0, line=0.98)],
    )
    if subtitle is not None:
        add_text(
            slide,
            MX,
            1.13,
            CW,
            0.38,
            [p(subtitle, sz=11.5, c=MUTED, after=0, line=1.05)],
        )


def footer(slide, number: int):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(MX),
        Inches(7.08),
        Inches(W - MX),
        Inches(7.08),
    )
    line.line.color.rgb = BORDER
    line.line.width = Pt(0.75)
    add_text(
        slide,
        MX,
        7.13,
        9.0,
        0.2,
        [
            p(
                "Clinical Report Assistant  \u00b7  synthetic-data capstone",
                sz=8,
                c=DIM,
                after=0,
            )
        ],
    )
    add_text(
        slide,
        W - MX - 0.8,
        7.13,
        0.8,
        0.2,
        [p(f"{number:02d}", sz=8, b=True, c=DIM, align=PP_ALIGN.RIGHT, after=0)],
    )


def section_label(slide, x, y, text, *, color=DIM, w=6.0):
    add_text(
        slide,
        x,
        y,
        w,
        0.2,
        [p(text.upper(), sz=8, b=True, c=color, after=0)],
    )


# ---------------------------------------------------------------------------
# Slide 1 - Title
# ---------------------------------------------------------------------------


def slide_title(prs):
    slide = base_slide(prs)

    # subtle corner frame
    rect(
        slide,
        8.35,
        0.95,
        4.45,
        5.45,
        fill=None,
        line=RGBColor(0x1E, 0x2C, 0x40),
        radius=0.04,
        line_w=1.0,
    )
    rect(
        slide,
        8.52,
        1.12,
        4.1,
        5.1,
        fill=None,
        line=RGBColor(0x18, 0x24, 0x36),
        radius=0.04,
        line_w=0.75,
        dash=MSO_LINE_DASH_STYLE.DASH,
    )

    add_text(
        slide,
        MX,
        1.12,
        7.0,
        0.3,
        [p("CAPSTONE PROJECT  \u00b7  5-MEMBER TEAM", sz=10, b=True, c=TEAL, after=0)],
    )
    add_text(
        slide,
        MX,
        1.48,
        7.4,
        1.75,
        [
            p("Clinical Report", sz=43, b=True, c=TEXT, after=0, line=0.95),
            p("Assistant", sz=43, b=True, c=TEAL, after=0, line=0.95),
        ],
    )
    add_text(
        slide,
        MX,
        3.38,
        7.25,
        1.05,
        [
            p(
                "A source-linked, human-in-the-loop pipeline that turns synthetic "
                "report PDFs into validated findings, cited guidance, and "
                "reviewable draft suggestions.",
                sz=13.5,
                c=MUTED,
                after=0,
                line=1.2,
            )
        ],
    )

    add_text(
        slide,
        MX,
        4.66,
        7.2,
        0.3,
        [
            rp(
                [
                    ("3 AI agents", {"sz": 11, "b": True, "c": TEXT}),
                    ("   \u00b7   1 orchestrated flow   \u00b7   ", {"sz": 11, "c": DIM}),
                    ("1 explicit human decision", {"sz": 11, "b": True, "c": AMBER}),
                ],
                after=0,
            )
        ],
    )

    chips = [
        ("Python 3.12", 1.52),
        ("Streamlit", 1.28),
        ("CrewAI Flow", 1.45),
        ("Gemini", 1.02),
        ("LangChain", 1.28),
        ("ChromaDB", 1.25),
        ("SQLite", 0.95),
        ("Pydantic", 1.15),
    ]
    x = MX
    for label, width in chips[:4]:
        chip(slide, x, 5.18, width, 0.32, label, sz=8.5)
        x += width + 0.12
    x = MX
    for label, width in chips[4:]:
        chip(slide, x, 5.6, width, 0.32, label, sz=8.5)
        x += width + 0.12

    add_text(
        slide,
        MX,
        6.62,
        7.3,
        0.3,
        [
            p(
                "Educational capstone  \u00b7  synthetic data only  \u00b7  not a medical device",
                sz=8.5,
                c=DIM,
                after=0,
            )
        ],
    )

    # right-hand pipeline graphic
    add_text(
        slide,
        8.85,
        1.3,
        3.7,
        0.22,
        [p("PIPELINE AT A GLANCE", sz=8, b=True, c=DIM, after=0)],
    )
    nodes = [
        ("Report PDF", "in-memory validation", TEAL),
        ("Validated facts", "page + excerpt evidence", TEAL),
        ("Summary + guidance", "bounded AI, local retrieval", BLUE),
        ("Draft suggestions", "citations or abstention", PURPLE),
        ("Human decision", "approve / reject exact draft", AMBER),
    ]
    y = 1.62
    for idx, (title, sub, color) in enumerate(nodes):
        rect(slide, 8.85, y, 3.65, 0.66, fill=CARD, line=BORDER, radius=0.14)
        rect(slide, 8.85, y, 0.055, 0.66, fill=color, line=None, shape=MSO_SHAPE.RECTANGLE)
        add_text(
            slide,
            9.02,
            y + 0.08,
            3.3,
            0.22,
            [p(title, sz=10.5, b=True, c=TEXT, after=0)],
        )
        add_text(
            slide,
            9.02,
            y + 0.33,
            3.3,
            0.2,
            [p(sub, sz=8.5, c=MUTED, after=0)],
        )
        if idx < len(nodes) - 1:
            down_arrow(slide, 10.67, y + 0.68, 0.16, color=DIM, w=0.13)
        y += 0.84

    slide.notes_slide.notes_text_frame.text = (
        "Presenter 5 opens. Introduce the project in one sentence: a local capstone "
        "that demonstrates safe PDF ingestion, source-linked analysis, guideline "
        "retrieval, and human review. Point to the three-agent claim, then to the "
        "right-hand pipeline: the system never jumps from PDF to advice. It stops "
        "at validated facts, then at a human decision. State clearly that this is "
        "a synthetic-data educational demo, not a medical device."
    )
    return slide


# ---------------------------------------------------------------------------
# Slide 2 - Problem
# ---------------------------------------------------------------------------


def slide_problem(prs):
    slide = base_slide(prs)
    header(
        slide,
        "The problem",
        "Clinical report review is slow and hard to trust",
        "A single PDF mixes measurements, units, reference ranges, and narrative "
        "context across pages. Manual review struggles to keep all four linked.",
    )

    items = [
        (
            "01",
            "Manual cross-checking",
            ("Values, units, and ranges are compared by eye, while context is "
            "scattered across pages and visits."),
            TEAL,
        ),
        (
            "02",
            "Fluent but unsupported AI",
            ("A summary can sound correct yet invent a value or cite a page that "
            "never contained it."),
            RED,
        ),
        (
            "03",
            "Arithmetic is not language",
            ("\u201cWithin\u201d, \u201cabove\u201d, or \u201cbelow\u201d must be computed from parsed "
            "numbers and an explicit supplied range."),
            BLUE,
        ),
        (
            "04",
            "No audit trail",
            ("Without immutable drafts and recorded decisions, reviewers cannot "
            "show what was approved \u2014 or what was reused."),
            PURPLE,
        ),
    ]
    positions = [(MX, 1.66), (6.72, 1.66), (MX, 3.28), (6.72, 3.28)]
    for (num, title, body, accent), (x, y) in zip(items, positions):
        rect(slide, x, y, 6.06, 1.42, fill=CARD, line=BORDER, radius=0.06)
        rect(slide, x, y, 0.05, 1.42, fill=accent, line=None, shape=MSO_SHAPE.RECTANGLE)
        add_text(
            slide,
            x + 0.22,
            y + 0.16,
            0.55,
            0.3,
            [p(num, sz=12, b=True, c=accent, after=0)],
        )
        add_text(
            slide,
            x + 0.78,
            y + 0.15,
            5.05,
            0.3,
            [p(title, sz=14, b=True, c=TEXT, after=0)],
        )
        add_text(
            slide,
            x + 0.78,
            y + 0.51,
            5.02,
            0.8,
            [p(body, sz=10.5, c=MUTED, after=0, line=1.12)],
        )

    rect(slide, MX, 5.02, CW, 1.02, fill=RGBColor(0x2A, 0x20, 0x0E), line=RGBColor(0x6B, 0x4F, 0x18), radius=0.06)
    rect(slide, MX, 5.02, 0.06, 1.02, fill=AMBER, line=None, shape=MSO_SHAPE.RECTANGLE)
    add_text(
        slide,
        MX + 0.28,
        5.18,
        2.6,
        0.6,
        [p("THE FAILURE MODE", sz=8.5, b=True, c=AMBER, after=2), p("we designed against", sz=8.5, c=RGBColor(0xD6, 0xB9, 0x7A), after=0)],
    )
    add_text(
        slide,
        MX + 2.7,
        5.22,
        9.2,
        0.6,
        [
            p(
                "A confident sentence with no traceable source \u2014 and no record "
                "of who accepted it.",
                sz=15.5,
                b=True,
                c=TEXT,
                after=0,
                line=1.02,
            )
        ],
    )

    add_text(
        slide,
        MX,
        6.42,
        CW,
        0.3,
        [
            p(
                "The design response: make evidence, arithmetic, and approval "
                "first-class parts of the workflow.",
                sz=9.5,
                c=DIM,
                after=0,
            )
        ],
    )

    slide.notes_slide.notes_text_frame.text = (
        "Presenter 1. Frame the problem as a chain of small failures, not one big "
        "one. Manual review is tiring and inconsistent; an LLM can produce a "
        "plausible sentence that is not supported by the document; numeric "
        "judgements are arithmetic and should be deterministic; and without an "
        "audit trail nobody can reconstruct the decision. End on the amber line: "
        "the dangerous output is confidence without provenance. This motivates "
        "the five design principles on the next slide."
    )
    return slide


# ---------------------------------------------------------------------------
# Slide 3 - Solution
# ---------------------------------------------------------------------------


def slide_solution(prs):
    slide = base_slide(prs)
    header(
        slide,
        "The solution",
        "Evidence first, AI bounded, humans in control",
        "A review assistant that assembles validated facts and cited options, "
        "then stops at an explicit human decision.",
    )

    steps = [
        ("1", "Ingest & validate", "PDF only, in memory"),
        ("2", "Analyze & compare", "facts + deterministic arithmetic"),
        ("3", "Summarize & retrieve", "one AI branch, one local branch"),
        ("4", "Recommend & review", "cite, abstain, or ask a human"),
    ]
    x = MX
    widths = [2.9, 2.95, 3.05, 2.95]
    for idx, ((num, title, sub), width) in enumerate(zip(steps, widths)):
        fill = RGBColor(0x10, 0x27, 0x2B) if idx < 2 else RGBColor(0x13, 0x1F, 0x35)
        rect(slide, x, 1.66, width, 0.78, fill=fill, line=BORDER, radius=0.12)
        add_text(slide, x + 0.14, 1.76, 0.28, 0.25, [p(num, sz=11, b=True, c=TEAL, after=0)])
        add_text(slide, x + 0.42, 1.73, width - 0.5, 0.24, [p(title, sz=10.5, b=True, c=TEXT, after=0)])
        add_text(slide, x + 0.42, 1.99, width - 0.5, 0.24, [p(sub, sz=8.5, c=MUTED, after=0)])
        if idx < 3:
            add_text(slide, x + width + 0.02, 1.9, 0.22, 0.2, [p("\u203a", sz=18, b=True, c=DIM, align=PP_ALIGN.CENTER, after=0)])
        x += width + 0.25

    principles = [
        (
            "Source-bound",
            ("Every finding keeps a page number and exact excerpt. Every suggestion "
            "quotes both patient and guideline evidence."),
            "page 4 + exact quote, every time",
            TEAL,
        ),
        (
            "Deterministic core",
            ("Range comparisons, repeat grouping, IDs, hashes, and compatibility "
            "checks run in code \u2014 not in the model."),
            "13.2 against a supplied range \u2192 code decides",
            BLUE,
        ),
        (
            "Bounded AI",
            ("Gemini returns structured data or finding IDs inside strict schemas; "
            "it never writes unchecked facts into the UI."),
            "IDs in, validated facts out",
            PURPLE,
        ),
        (
            "Abstain by default",
            ("Absent, unrelated, contradictory, or insufficient guidance produces "
            "no suggestion."),
            "no supporting passage \u2192 no suggestion",
            AMBER,
        ),
        (
            "Local & reusable",
            ("Embeddings, Chroma, SQLite, and review records stay local. Compatible "
            "results skip new Gemini calls."),
            "same hash + settings \u2192 reuse, zero calls",
            GREEN,
        ),
    ]
    x = MX
    for title, body, example, accent in principles:
        rect(slide, x, 2.72, 2.35, 2.38, fill=CARD, line=BORDER, radius=0.06)
        rect(slide, x, 2.72, 2.35, 0.055, fill=accent, line=None, shape=MSO_SHAPE.RECTANGLE)
        add_text(
            slide,
            x + 0.17,
            2.9,
            2.0,
            0.5,
            [p(title, sz=12, b=True, c=TEXT, after=0, line=1.0)],
        )
        add_text(
            slide,
            x + 0.17,
            3.38,
            2.0,
            1.05,
            [p(body, sz=9.2, c=MUTED, after=0, line=1.13)],
        )
        add_text(
            slide,
            x + 0.17,
            4.56,
            2.0,
            0.42,
            [p("\u2192 " + example, sz=7.9, c=accent, after=0, line=1.0)],
        )
        x += 2.47

    rect(slide, MX, 5.35, CW, 1.22, fill=CARD_ALT, line=BORDER, radius=0.06)
    outcomes = [
        ("Traceable", "every claim has a page and an excerpt", TEAL),
        ("Repeatable", "same validated inputs, same deterministic output", BLUE),
        ("Reviewable", "approval binds to one immutable draft", AMBER),
    ]
    x = MX + 0.3
    for idx, (title, sub, color) in enumerate(outcomes):
        add_text(slide, x, 5.52, 3.6, 0.28, [p(title, sz=12.5, b=True, c=color, after=0)])
        add_text(slide, x, 5.82, 3.55, 0.45, [p(sub, sz=9.2, c=MUTED, after=0, line=1.05)])
        x += 4.0

    slide.notes_slide.notes_text_frame.text = (
        "Presenter 3 introduces the solution before the architecture. Read the "
        "four-step ribbon left to right, then slow down on the five principles. "
        "The important distinction is that AI is used where language is useful "
        "and constrained where correctness matters. 'Abstain by default' is a "
        "feature, not a failure: the system would rather return no suggestion "
        "than cite guidance it cannot support. The bottom row previews the "
        "qualities the rest of the deck will prove."
    )
    return slide


# ---------------------------------------------------------------------------
# Slide 4 - Team ownership
# ---------------------------------------------------------------------------


def slide_team(prs):
    slide = base_slide(prs)
    header(
        slide,
        "Ownership",
        "Five workstreams, one integrated product",
        "Each presenter owns a vertical slice; integration is a shared contract, "
        "not a hand-off at the end.",
    )

    people = [
        (
            "1",
            "Ingestion & Report Analysis",
            ("Validate the PDF, extract text page by page, and convert findings "
            "into source-checked facts with deterministic range comparisons."),
            "validated ReportAnalysisResult",
            "Presenter 1",
            TEAL,
        ),
        (
            "2",
            "Guideline RAG",
            ("Chunk guideline PDFs, embed them locally, persist them in Chroma, "
            "and return passages with document and page references."),
            "page-linked passages + index fingerprint",
            "Presenter 2",
            PURPLE,
        ),
        (
            "3",
            "Summary & Presentation",
            ("Reduce validated analysis to a concise summary, preserve evidence, "
            "group repeat measurements, and surface uncertainty."),
            "validated ReportSummaryResult",
            "Presenter 3",
            BLUE,
        ),
        (
            "4",
            "Recommendation & Review",
            ("Draft cited review suggestions or abstain; bind human approval to "
            "one exact draft and version every edit."),
            "immutable draft + review decision",
            "Presenter 4",
            AMBER,
        ),
        (
            "5",
            "Orchestration & Persistence",
            ("Sequence the agents, overlap summary and retrieval, persist immutable "
            "results, and reuse them without new Gemini calls."),
            "flow routes + reusable saved stages",
            "Presenter 5",
            GREEN,
        ),
    ]

    x = MX
    for num, title, body, output, owner, color in people:
        rect(slide, x, 1.68, 2.35, 3.95, fill=CARD, line=BORDER, radius=0.05)
        rect(slide, x, 1.68, 2.35, 0.06, fill=color, line=None, shape=MSO_SHAPE.RECTANGLE)
        add_text(slide, x + 0.17, 1.88, 0.5, 0.3, [p(num, sz=17, b=True, c=color, after=0)])
        add_text(
            slide,
            x + 0.17,
            2.28,
            2.02,
            0.85,
            [p(title, sz=12, b=True, c=TEXT, after=0, line=1.02)],
        )
        add_text(
            slide,
            x + 0.17,
            3.2,
            2.02,
            1.2,
            [p(body, sz=9.2, c=MUTED, after=0, line=1.15)],
        )
        line = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            Inches(x + 0.17),
            Inches(4.5),
            Inches(x + 2.18),
            Inches(4.5),
        )
        line.line.color.rgb = BORDER
        line.line.width = Pt(0.75)
        add_text(
            slide,
            x + 0.17,
            4.6,
            2.02,
            0.2,
            [p("OUTPUT", sz=7.2, b=True, c=color, after=0)],
        )
        add_text(
            slide,
            x + 0.17,
            4.78,
            2.02,
            0.4,
            [p(output, sz=8.1, c=MUTED, after=0, line=1.0)],
        )
        add_text(
            slide,
            x + 0.17,
            5.18,
            2.02,
            0.25,
            [p(owner.upper(), sz=8, b=True, c=color, after=0)],
        )
        x += 2.47

    rect(slide, MX, 5.86, CW, 0.86, fill=CARD_ALT, line=BORDER, radius=0.06)
    add_text(
        slide,
        MX + 0.25,
        5.99,
        CW - 0.5,
        0.6,
        [
            rp(
                [
                    ("Module boundaries are enforced in code.  ", {"sz": 10, "b": True, "c": TEAL}),
                    (
                        ("The UI never owns extraction rules; workers never share "
                        "open SQLite or Chroma connections; agents never persist "
                        "their own output without validation."),
                        {"sz": 10, "c": MUTED},
                    ),
                ],
                after=0,
                line=1.1,
            )
        ],
    )

    slide.notes_slide.notes_text_frame.text = (
        "All five presenters stand or speak in turn: one sentence each on the "
        "workstream they own. Emphasise that these are not five separate "
        "projects. The contracts between modules are tested: the presentation "
        "layer calls processing code, agents return structured models, and the "
        "orchestrator owns persistence and reuse. The bottom line is the "
        "architectural rule that keeps the demo safe and debuggable."
    )
    return slide


# ---------------------------------------------------------------------------
# Slide 5 - Architecture
# ---------------------------------------------------------------------------


def slide_architecture(prs):
    slide = base_slide(prs)
    header(
        slide,
        "Architecture",
        "Overall architecture",
        "One presentation layer, one orchestrator, exactly three AI agents, and "
        "deterministic support around them.",
    )

    # Presentation layer
    rect(slide, MX, 1.6, CW, 0.62, fill=RGBColor(0x10, 0x22, 0x38), line=RGBColor(0x2B, 0x4A, 0x6E), radius=0.10)
    add_text(
        slide,
        MX + 0.22,
        1.72,
        1.75,
        0.3,
        [p("Streamlit UI", sz=11.5, b=True, c=BLUE, after=0)],
    )
    add_text(
        slide,
        MX + 1.98,
        1.74,
        10.0,
        0.3,
        [
            p(
                "upload  \u00b7  explicit run controls  \u00b7  compact results  \u00b7  "
                "evidence expanders  \u00b7  approve / reject",
                sz=9.5,
                c=MUTED,
                after=0,
            )
        ],
    )
    down_arrow(slide, 6.67, 2.24, 0.15, color=BLUE, w=0.15)

    # Orchestration
    rect(slide, MX, 2.41, CW, 0.62, fill=RGBColor(0x10, 0x22, 0x38), line=RGBColor(0x2B, 0x4A, 0x6E), radius=0.10)
    add_text(
        slide,
        MX + 0.22,
        2.53,
        1.9,
        0.3,
        [p("CrewAI Flow", sz=11.5, b=True, c=BLUE, after=0)],
    )
    add_text(
        slide,
        MX + 2.0,
        2.55,
        9.9,
        0.3,
        [
            p(
                "resolve analysis  \u2192  fork two workers  \u2192  join  \u2192  "
                "re-check index  \u2192  recommend  \u2192  route outcome",
                sz=9.5,
                c=MUTED,
                after=0,
            )
        ],
    )
    down_arrow(slide, 3.0, 3.05, 0.17, color=BLUE, w=0.15)
    down_arrow(slide, 10.7, 3.05, 0.17, color=PURPLE, w=0.15)

    # Agent container
    rect(
        slide,
        0.55,
        3.24,
        7.98,
        2.1,
        fill=CARD_ALT,
        line=RGBColor(0x2B, 0x4A, 0x6E),
        radius=0.04,
        dash=MSO_LINE_DASH_STYLE.DASH,
    )
    add_text(
        slide,
        0.72,
        3.3,
        4.8,
        0.2,
        [p("AGENT LAYER  \u00b7  CrewAI + native Gemini", sz=7.8, b=True, c=BLUE, after=0)],
    )
    chip(
        slide,
        5.92,
        3.27,
        2.45,
        0.24,
        "Gemini API  \u00b7  external",
        fill=RGBColor(0x2A, 0x20, 0x0E),
        line=RGBColor(0x6B, 0x4F, 0x18),
        color=AMBER,
        sz=7.5,
        bold=True,
    )

    agents = [
        (
            "Report Analysis",
            "extract  \u00b7  validate  \u00b7  compare",
            "findings + narrative in a strict schema",
            "structured output",
            TEAL,
        ),
        (
            "Summary",
            "select finding IDs",
            "deterministic assembly, evidence kept",
            "IDs only",
            TEAL,
        ),
        (
            "Recommendation",
            "cite or abstain",
            "patient + guideline evidence",
            "cited draft",
            TEAL,
        ),
    ]
    x = 0.73
    for title, sub, detail, gemini, color in agents:
        rect(slide, x, 3.6, 2.43, 1.42, fill=CARD, line=BORDER, radius=0.08)
        rect(slide, x, 3.6, 2.43, 0.045, fill=color, line=None, shape=MSO_SHAPE.RECTANGLE)
        add_text(slide, x + 0.14, 3.73, 2.12, 0.26, [p(title, sz=10.5, b=True, c=TEXT, after=0)])
        add_text(slide, x + 0.14, 4.02, 2.12, 0.34, [p(sub, sz=8.4, c=MUTED, after=0, line=1.02)])
        add_text(slide, x + 0.14, 4.34, 2.12, 0.26, [p(detail, sz=7.7, c=MUTED, after=0, line=1.0)])
        add_text(
            slide,
            x + 0.14,
            4.66,
            2.12,
            0.25,
            [p("Gemini: " + gemini, sz=7.6, c=AMBER, after=0)],
        )
        x += 2.58

    # RAG
    rect(slide, 8.7, 3.24, 4.08, 2.1, fill=RGBColor(0x1A, 0x16, 0x2E), line=RGBColor(0x4A, 0x3C, 0x78), radius=0.04)
    add_text(slide, 8.88, 3.36, 3.7, 0.25, [p("Guideline KB  \u00b7  local RAG", sz=10.5, b=True, c=PURPLE, after=0)])
    add_text(
        slide,
        8.88,
        3.7,
        3.7,
        1.0,
        [
            p("page-aware chunks", sz=8.8, c=TEXT, after=2),
            p("MiniLM-L6-v2 embeddings  \u00b7  CPU", sz=8.8, c=TEXT, after=2),
            p("persistent Chroma collection", sz=8.8, c=TEXT, after=2),
            p("top-k passages with page + hash", sz=8.8, c=TEXT, after=0),
        ],
    )
    add_text(
        slide,
        8.88,
        4.68,
        3.65,
        0.55,
        [p("Read-only for the report workflow; report uploads never enter this store.", sz=7.8, c=MUTED, after=0, line=1.05)],
    )

    # Foundation
    section_label(slide, MX, 5.47, "Deterministic support", color=DIM)
    foundation = [
        ("pdf_processing", "in-memory PDF gates"),
        ("Pydantic boundaries", "validated evidence models"),
        ("SQLite store", "immutable results + review"),
        ("session_state", "stale-result clearing"),
        ("config", "env \u2192 secrets \u2192 .env"),
    ]
    x = MX
    for title, sub in foundation:
        rect(slide, x, 5.7, 2.33, 0.58, fill=CARD, line=BORDER, radius=0.12)
        add_text(slide, x + 0.12, 5.76, 2.1, 0.2, [p(title, sz=8.6, b=True, c=TEXT, after=0)])
        add_text(slide, x + 0.12, 5.96, 2.1, 0.2, [p(sub, sz=7.5, c=MUTED, after=0)])
        x += 2.45

    add_text(
        slide,
        MX,
        6.42,
        CW,
        0.24,
        [
            rp(
                [
                    ("Legend:  ", {"sz": 7.8, "b": True, "c": DIM}),
                    ("teal = AI agents   ", {"sz": 7.8, "c": TEAL}),
                    ("blue = application control   ", {"sz": 7.8, "c": BLUE}),
                    ("purple = local retrieval   ", {"sz": 7.8, "c": PURPLE}),
                    ("amber = external model   ", {"sz": 7.8, "c": AMBER}),
                    ("grey = deterministic support", {"sz": 7.8, "c": MUTED}),
                ],
                after=0,
            )
        ],
    )
    add_text(
        slide,
        MX,
        6.7,
        CW,
        0.24,
        [
            p(
                "Trust boundary: uploaded reports and retrieved passages are untrusted data. "
                "Validation runs before anything is displayed or reused.",
                sz=8.2,
                c=DIM,
                after=0,
            )
        ],
    )

    slide.notes_slide.notes_text_frame.text = (
        "Presenter 5. Walk the stack from top to bottom. The Streamlit layer only "
        "presents and collects explicit actions; it does not own extraction rules. "
        "CrewAI Flow is the only orchestrator. Inside the agent layer, exactly "
        "three agents call Gemini, and each call is schema-bounded. Guideline RAG "
        "is application-level retrieval, not an agent-invoked tool. The foundation "
        "row is deliberately deterministic: PDF gates, Pydantic, SQLite, session "
        "state, and configuration. Finish on the trust boundary: nothing from the "
        "PDF or the guideline store is trusted until it passes validation."
    )
    return slide


# ---------------------------------------------------------------------------
# Slide 6 - Workflow
# ---------------------------------------------------------------------------


def slide_workflow(prs):
    slide = base_slide(prs)
    header(
        slide,
        "Workflow",
        "End-to-end workflow: reuse when safe",
        "Every stage first looks for an exact compatible record; a hit skips the "
        "model call. A miss generates, validates, and persists.",
    )

    # Band 1
    band1 = [
        ("1", "Explicit run", "User selects Run remaining workflow"),
        ("2", "Resolve analysis", "reuse exact record, or call Report Analysis"),
        ("3", "Validate & persist", "source checks pass \u2192 immutable analysis"),
    ]
    x = MX
    for num, title, sub in band1:
        rect(slide, x, 1.64, 3.9, 0.9, fill=CARD, line=BORDER, radius=0.09)
        add_text(slide, x + 0.16, 1.74, 0.3, 0.25, [p(num, sz=10, b=True, c=TEAL, after=0)])
        add_text(slide, x + 0.46, 1.71, 3.2, 0.26, [p(title, sz=11.5, b=True, c=TEXT, after=0)])
        add_text(slide, x + 0.46, 1.99, 3.25, 0.4, [p(sub, sz=8.6, c=MUTED, after=0, line=1.05)])
        x += 4.16

    # connectors
    right_arrow_shape(slide, 4.47, 2.0, 0.22, 0.16)
    right_arrow_shape(slide, 8.63, 2.0, 0.22, 0.16)

    # Parallel band
    rect(
        slide,
        0.55,
        2.88,
        12.23,
        1.62,
        fill=RGBColor(0x10, 0x22, 0x38),
        line=RGBColor(0x2B, 0x4A, 0x6E),
        radius=0.05,
        dash=MSO_LINE_DASH_STYLE.DASH,
    )
    add_text(
        slide,
        0.72,
        2.94,
        6.0,
        0.2,
        [p("PARALLEL BRANCHES  \u00b7  max_workers = 2", sz=7.8, b=True, c=BLUE, after=0)],
    )
    add_text(
        slide,
        9.0,
        2.94,
        3.6,
        0.2,
        [p("one AI branch + one local branch", sz=7.8, c=MUTED, align=PP_ALIGN.RIGHT, after=0)],
    )

    rect(slide, 0.9, 3.18, 5.6, 1.05, fill=CARD, line=BORDER, radius=0.09)
    add_text(slide, 1.08, 3.29, 5.1, 0.27, [p("4a  Summary Agent", sz=11, b=True, c=TEAL, after=0)])
    add_text(slide, 1.08, 3.58, 5.1, 0.26, [p("reuse compatible summary, or generate \u00b7 validate \u00b7 persist", sz=8.6, c=MUTED, after=0, line=1.05)])
    add_text(slide, 1.08, 3.84, 5.1, 0.24, [p("every returned ID checked; facts rebuilt by code", sz=7.8, c=DIM, after=0)])

    rect(slide, 6.95, 3.18, 5.55, 1.05, fill=RGBColor(0x1A, 0x16, 0x2E), line=RGBColor(0x4A, 0x3C, 0x78), radius=0.09)
    add_text(slide, 7.13, 3.29, 5.1, 0.27, [p("4b  Guideline retrieval", sz=11, b=True, c=PURPLE, after=0)])
    add_text(slide, 7.13, 3.58, 5.1, 0.26, [p("queries from analysis \u00b7 passages + exact index fingerprint", sz=8.6, c=MUTED, after=0, line=1.05)])
    add_text(slide, 7.13, 3.84, 5.1, 0.24, [p("needs no summary; runs while the Summary Agent works", sz=7.8, c=DIM, after=0)])

    # flow between bands
    down_arrow(slide, 6.67, 2.57, 0.26, color=BLUE, w=0.14)
    down_arrow(slide, 2.4, 4.53, 0.19, color=AMBER, w=0.14)

    # Band 3
    band3 = [
        ("4", "Join + re-check", "both branches required; index must be unchanged"),
        ("5", "Recommendation", "reuse or generate a cited draft"),
        ("6", "Human decision", "approve / reject the exact draft"),
    ]
    x = MX
    for num, title, sub in band3:
        rect(slide, x, 4.76, 3.9, 0.9, fill=CARD, line=BORDER, radius=0.09)
        add_text(slide, x + 0.16, 4.86, 0.3, 0.25, [p(num, sz=10, b=True, c=AMBER, after=0)])
        add_text(slide, x + 0.46, 4.83, 3.2, 0.26, [p(title, sz=11.5, b=True, c=TEXT, after=0)])
        add_text(slide, x + 0.46, 5.11, 3.25, 0.4, [p(sub, sz=8.6, c=MUTED, after=0, line=1.05)])
        x += 4.16

    right_arrow_shape(slide, 4.47, 5.12, 0.22, 0.16)
    right_arrow_shape(slide, 8.63, 5.12, 0.22, 0.16)

    # Route strip
    add_text(slide, MX, 5.88, 2.4, 0.2, [p("RECORDED ROUTES", sz=7.8, b=True, c=DIM, after=0)])
    routes = ["reuse", "generate", "retrieve", "abstain", "await_review", "approved", "rejected", "failed"]
    x = MX
    for label in routes:
        chip(slide, x, 6.12, 1.42, 0.27, label, sz=7.6, color=MUTED)
        x += 1.5

    add_text(
        slide,
        MX,
        6.58,
        CW,
        0.3,
        [
            p(
                "A failed branch never deletes a successfully saved upstream stage. "
                "The join requires both branches; a changed index blocks recommendation reuse.",
                sz=8.4,
                c=DIM,
                after=0,
            )
        ],
    )

    slide.notes_slide.notes_text_frame.text = (
        "Presenter 5. Start with the explicit run: nothing calls Gemini on upload "
        "or page refresh. Analysis resolves first because both later branches "
        "depend on it. Then the flow forks into one Summary Agent branch and one "
        "local retrieval branch, capped at two workers. Retrieval uses only "
        "validated analysis, so it does not wait for the summary. The join "
        "requires both, then the flow re-checks the guideline fingerprint before "
        "recommendation. Point to the route strip: these are auditable decisions, "
        "not UI labels. Human review is the only place an approval can happen."
    )
    return slide


# ---------------------------------------------------------------------------
# Slide 7 - Presenter 1
# ---------------------------------------------------------------------------


def slide_p1(prs):
    slide = base_slide(prs)
    header(
        slide,
        "Presenter 1  \u00b7  ingestion + analysis",
        "Ingestion and Report Analysis Agent",
        "Turn an untrusted PDF into validated, page-linked facts \u2014 and let code do the arithmetic.",
    )

    # Left panel
    rect(slide, MX, 1.66, 4.28, 4.62, fill=CARD_ALT, line=BORDER, radius=0.05)
    add_text(slide, MX + 0.2, 1.8, 3.8, 0.28, [p("Ingestion gates", sz=13, b=True, c=TEAL, after=0)])
    add_text(slide, MX + 0.2, 2.08, 3.8, 0.22, [p("in memory, before any model call", sz=8.5, c=MUTED, after=0)])

    gates = [
        ("Size & type", "non-empty, \u2264 10 MiB, %PDF- header"),
        ("Parse safely", "pypdf must open it; encrypted files rejected"),
        ("Page extraction", "text extracted page by page, one-based"),
        ("No text?", "warning + OCR limitation; nothing invented"),
        ("Identity", "SHA-256 of bytes, never the filename"),
    ]
    y = 2.42
    for idx, (title, sub) in enumerate(gates, start=1):
        badge(slide, MX + 0.22, y + 0.03, 0.3, str(idx), color=TEAL)
        add_text(slide, MX + 0.64, y, 3.35, 0.22, [p(title, sz=9.8, b=True, c=TEXT, after=0)])
        add_text(slide, MX + 0.64, y + 0.21, 3.35, 0.24, [p(sub, sz=8.2, c=MUTED, after=0)])
        y += 0.62

    rect(slide, MX + 0.2, 5.62, 3.85, 0.48, fill=RGBColor(0x0D, 0x2A, 0x2C), line=RGBColor(0x1F, 0x5A, 0x55), radius=0.15)
    add_text(
        slide,
        MX + 0.32,
        5.7,
        3.6,
        0.3,
        [p("Original PDF bytes are never persisted.", sz=8.6, b=True, c=TEAL, after=0)],
    )

    # Right panel
    rect(slide, 5.0, 1.66, 7.78, 4.62, fill=CARD, line=BORDER, radius=0.05)
    add_text(slide, 5.2, 1.8, 6.8, 0.28, [p("Report Analysis Agent", sz=13, b=True, c=TEAL, after=0)])
    add_text(slide, 5.2, 2.08, 7.2, 0.22, [p("bounded extraction, then deterministic validation", sz=8.5, c=MUTED, after=0)])

    steps = [
        (
            "1",
            "Bounded extraction",
            "CrewAI + Gemini returns findings and narrative facts in a strict schema.",
        ),
        (
            "2",
            "Source validation",
            "The cited page must contain the quoted excerpt; values, units, and reference text are checked.",
        ),
        (
            "3",
            "Deterministic comparison",
            "Explicit ranges are parsed with Decimal \u2192 within / below / above / not assessed.",
        ),
        (
            "4",
            "Separated outputs",
            "Patient measurements, narrative facts, labeled training controls, and review warnings stay distinct.",
        ),
    ]
    y = 2.42
    for idx, (num, title, body) in enumerate(steps):
        add_text(slide, 5.2, y, 0.32, 0.26, [p(num, sz=10, b=True, c=BLUE, after=0)])
        add_text(slide, 5.52, y - 0.02, 6.9, 0.26, [p(title, sz=10.5, b=True, c=TEXT, after=0)])
        add_text(slide, 5.52, y + 0.25, 6.95, 0.5, [p(body, sz=8.8, c=MUTED, after=0, line=1.08)])
        if idx < 3:
            line = slide.shapes.add_connector(
                MSO_CONNECTOR.STRAIGHT,
                Inches(5.36),
                Inches(y + 0.72),
                Inches(5.36),
                Inches(y + 0.9),
            )
            line.line.color.rgb = BORDER
            line.line.width = Pt(1.0)
        y += 0.9

    # Bottom strip
    rect(slide, MX, 6.44, CW, 0.46, fill=RGBColor(0x2A, 0x20, 0x0E), line=RGBColor(0x6B, 0x4F, 0x18), radius=0.2)
    add_text(
        slide,
        MX + 0.22,
        6.54,
        CW - 0.44,
        0.3,
        [
            rp(
                [
                    ("Never guessed:  ", {"sz": 8.4, "b": True, "c": AMBER}),
                    (
                        ("missing range \u2192 not_assessed  \u00b7  unit mismatch \u2192 review  "
                        "\u00b7  ambiguous context \u2192 review  \u00b7  repeated analyte \u2192 needs a discriminator"),
                        {"sz": 8.4, "c": TEXT},
                    ),
                ],
                after=0,
            )
        ],
    )

    slide.notes_slide.notes_text_frame.text = (
        "Presenter 1. The left panel is a set of gates, not a single 'upload' step: "
        "size, header, parseability, encryption, page extraction, and identity. "
        "A scanned PDF gets a visible warning; the system never invents text. On "
        "the right, the agent first proposes structured findings. Then code checks "
        "that every cited excerpt exists on the cited page, validates units and "
        "reference text, and performs the numeric comparison with Decimal. The "
        "fourth step matters: patient measurements, narrative facts, training "
        "controls, and warnings are separate fields, so a control value can never "
        "silently become a patient result."
    )
    return slide


# ---------------------------------------------------------------------------
# Slide 8 - Presenter 2
# ---------------------------------------------------------------------------


def slide_p2(prs):
    slide = base_slide(prs)
    header(
        slide,
        "Presenter 2  \u00b7  guideline RAG",
        "Guideline RAG: from PDF pages to cited passages",
        "Reference guidance is chunked, embedded, and stored locally so every "
        "retrieved passage still knows its document and page.",
    )

    pipeline = [
        ("1", "Guideline PDF", "selectable text only"),
        ("2", "Page-aware chunks", "800 chars \u00b7 120 overlap"),
        ("3", "Local embeddings", "MiniLM-L6-v2 \u00b7 384-d \u00b7 CPU"),
        ("4", "ChromaDB", "persistent local collection"),
        ("5", "Top-k passages", "text + document + page + hash"),
    ]
    x = MX
    for num, title, sub in pipeline:
        rect(slide, x, 1.64, 2.22, 1.18, fill=CARD, line=BORDER, radius=0.08)
        add_text(slide, x + 0.14, 1.73, 0.28, 0.22, [p(num, sz=9, b=True, c=PURPLE, after=0)])
        add_text(slide, x + 0.14, 1.95, 1.95, 0.34, [p(title, sz=10.3, b=True, c=TEXT, after=0, line=1.0)])
        add_text(slide, x + 0.14, 2.34, 1.95, 0.34, [p(sub, sz=8.1, c=MUTED, after=0, line=1.02)])
        x += 2.47

    for cx in (2.81, 5.28, 7.75, 10.22):
        right_arrow_shape(slide, cx, 2.19, 0.2, 0.16, color=PURPLE)

    cards = [
        (
            "Deterministic chunk ID",
            "SHA-256(document hash : page : chunk index : text). Re-indexing the same PDF is detected by document hash and skipped.",
            "same bytes \u2192 same chunks",
            PURPLE,
        ),
        (
            "Metadata contract",
            "Every chunk carries document name and hash, one-based page, chunk ID, embedding model, and splitter settings.",
            "carried with every passage",
            BLUE,
        ),
        (
            "Index fingerprint",
            "Manifest of document hashes + configuration \u2192 one SHA-256. Recommendation reuse is invalidated when the index changes.",
            "bound into recommendation reuse",
            AMBER,
        ),
    ]
    x = MX
    for title, body, tag, accent in cards:
        rect(slide, x, 3.13, 3.95, 1.42, fill=CARD, line=BORDER, radius=0.06)
        rect(slide, x, 3.13, 0.05, 1.42, fill=accent, line=None, shape=MSO_SHAPE.RECTANGLE)
        add_text(slide, x + 0.18, 3.25, 3.55, 0.26, [p(title, sz=10.5, b=True, c=TEXT, after=0)])
        add_text(slide, x + 0.18, 3.55, 3.58, 0.62, [p(body, sz=8.5, c=MUTED, after=0, line=1.08)])
        add_text(slide, x + 0.18, 4.22, 3.58, 0.22, [p(tag.upper(), sz=7.1, b=True, c=accent, after=0)])
        x += 4.14

    # Passage mock
    rect(slide, MX, 4.82, 6.03, 1.55, fill=RGBColor(0x1A, 0x16, 0x2E), line=RGBColor(0x4A, 0x3C, 0x78), radius=0.06)
    add_text(slide, MX + 0.2, 4.93, 5.5, 0.2, [p("RETRIEVED PASSAGE  \u00b7  SHAPE", sz=7.8, b=True, c=PURPLE, after=0)])
    add_text(
        slide,
        MX + 0.2,
        5.17,
        5.65,
        0.35,
        [
            p(
                "synthetic_capstone_review_guidance.pdf  \u00b7  page 3",
                sz=9,
                b=True,
                c=TEXT,
                after=0,
            )
        ],
    )
    add_text(
        slide,
        MX + 0.2,
        5.48,
        5.65,
        0.35,
        [
            p(
                "\u201cExact passage text is retained for citation.\u201d",
                sz=9,
                c=RGBColor(0xD6, 0xCE, 0xF0),
                i=True,
                after=0,
            )
        ],
    )
    add_text(
        slide,
        MX + 0.2,
        5.91,
        5.65,
        0.26,
        [p("chunk ID  \u00b7  document hash  \u00b7  one-based page", sz=7.8, c=MUTED, after=0)],
    )

    # Guardrails
    rect(slide, 6.78, 4.82, 6.0, 1.55, fill=CARD_ALT, line=BORDER, radius=0.06)
    add_text(slide, 6.96, 4.93, 5.6, 0.2, [p("GUARDRAILS", sz=7.8, b=True, c=AMBER, after=0)])
    add_text(
        slide,
        6.96,
        5.17,
        5.6,
        1.05,
        [
            bullet("Similarity ranks relevance; it is not a confidence score.", sz=8.8, after=4),
            bullet("Incompatible or unmanifested indexes are rejected.", sz=8.8, after=4),
            bullet("Report uploads never enter the guideline store.", sz=8.8, after=4),
            bullet("Empty index \u2192 no passages available for a suggestion.", sz=8.8, after=0),
        ],
    )

    add_text(
        slide,
        MX,
        6.55,
        CW,
        0.3,
        [
            p(
                "Retrieval is application-level RAG: deterministic queries, bounded top-k, "
                "and no agent-invoked vector search.",
                sz=8.4,
                c=DIM,
                after=0,
            )
        ],
    )

    slide.notes_slide.notes_text_frame.text = (
        "Presenter 2. Follow one guideline PDF through the five stages. Chunking "
        "is page-aware, so a passage can never drift to another page. Embeddings "
        "run locally with MiniLM on CPU, normalized, and Chroma persists under the "
        "ignored data directory. Every chunk keeps a deterministic ID and its full "
        "source metadata. The manifest and fingerprint are what make reuse safe: "
        "if the index changes, old recommendations are no longer compatible. "
        "Close with the guardrail that similarity is a ranking mechanism, not a "
        "clinical confidence score."
    )
    return slide


# ---------------------------------------------------------------------------
# Slide 9 - Presenter 3
# ---------------------------------------------------------------------------


def slide_p3(prs):
    slide = base_slide(prs)
    header(
        slide,
        "Presenter 3  \u00b7  summary + evidence",
        "Summary Agent: select less, verify more",
        "Gemini chooses finding IDs; deterministic code rebuilds every fact, "
        "page, repeat group, and uncertainty.",
    )

    # Left
    rect(slide, MX, 1.64, 5.9, 2.12, fill=CARD, line=BORDER, radius=0.05)
    add_text(slide, MX + 0.2, 1.76, 5.4, 0.25, [p("Gemini returns IDs only", sz=12, b=True, c=TEAL, after=0)])
    add_text(
        slide,
        MX + 0.2,
        2.07,
        5.45,
        1.55,
        [
            bullet("overview_finding_ids", sz=8.9, after=4),
            bullet("important_finding_ids", sz=8.9, after=4),
            bullet("repeat_result_groups (ordering hint)", sz=8.9, after=4),
            bullet("narrative_finding_ids", sz=8.9, after=4),
            bullet("uncertainty_finding_ids (ordering hint)", sz=8.9, after=0),
        ],
    )
    add_text(
        slide,
        MX + 0.2,
        3.47,
        5.4,
        0.22,
        [p("Bounded lists, strict schema, no free-text facts.", sz=8, c=AMBER, after=0)],
    )

    # Right
    rect(slide, 6.65, 1.64, 6.13, 2.12, fill=CARD, line=BORDER, radius=0.05)
    add_text(slide, 6.85, 1.76, 5.6, 0.25, [p("Code rebuilds the summary", sz=12, b=True, c=BLUE, after=0)])
    add_text(
        slide,
        6.85,
        2.07,
        5.65,
        1.55,
        [
            bullet("Numeric values, units, and reference ranges", sz=8.9, bc=BLUE, after=4),
            bullet("Repeat sequences and supported differences", sz=8.9, bc=BLUE, after=4),
            bullet("Narrative statements and attribution", sz=8.9, bc=BLUE, after=4),
            bullet("Pages, excerpts, and applicability evidence", sz=8.9, bc=BLUE, after=4),
            bullet("Mandatory uncertainties derived from validated analysis", sz=8.9, bc=BLUE, after=0),
        ],
    )

    # Section cards
    add_text(slide, MX, 3.92, 5.0, 0.2, [p("SUMMARY SECTIONS", sz=7.8, b=True, c=DIM, after=0)])
    sections = [
        ("Overview", "one validated orientation item", TEAL),
        ("Important findings", "selected, evidence-linked", BLUE),
        ("Repeat measurements", "normalized analyte + matching unit", PURPLE),
        ("Narrative context", "symptoms, history, gaps", GREEN),
        ("Uncertainties", "derived, never dropped", AMBER),
    ]
    x = MX
    for title, sub, color in sections:
        rect(slide, x, 4.16, 2.35, 1.18, fill=CARD_ALT, line=BORDER, radius=0.08)
        rect(slide, x, 4.16, 2.35, 0.045, fill=color, line=None, shape=MSO_SHAPE.RECTANGLE)
        add_text(slide, x + 0.15, 4.28, 2.05, 0.42, [p(title, sz=9.8, b=True, c=TEXT, after=0, line=1.0)])
        add_text(slide, x + 0.15, 4.76, 2.05, 0.45, [p(sub, sz=8.1, c=MUTED, after=0, line=1.03)])
        x += 2.47

    # Callouts
    rect(slide, MX, 5.58, 6.0, 1.02, fill=RGBColor(0x1A, 0x16, 0x2E), line=RGBColor(0x4A, 0x3C, 0x78), radius=0.06)
    add_text(slide, MX + 0.18, 5.69, 5.6, 0.22, [p("REPEAT RULE", sz=7.8, b=True, c=PURPLE, after=0)])
    add_text(
        slide,
        MX + 0.18,
        5.94,
        5.65,
        0.55,
        [p("Group only patient measurements with normalized analyte identity and matching units. Never invent chronology.", sz=8.5, c=TEXT, after=0, line=1.05)],
    )

    rect(slide, 6.75, 5.58, 6.03, 1.02, fill=CARD, line=BORDER, radius=0.06)
    add_text(slide, 6.93, 5.69, 5.6, 0.22, [p("UI RULE", sz=7.8, b=True, c=TEAL, after=0)])
    add_text(
        slide,
        6.93,
        5.94,
        5.65,
        0.55,
        [p("Compact rows by default; full items and evidence stay in expanders; warnings remain visible.", sz=8.5, c=TEXT, after=0, line=1.05)],
    )

    add_text(
        slide,
        MX,
        6.72,
        CW,
        0.25,
        [p("Absent optional context does not create uncertainty; mandatory uncertainty is derived from the analysis.", sz=8.2, c=DIM, after=0)],
    )

    slide.notes_slide.notes_text_frame.text = (
        "Presenter 3. The Summary Agent is deliberately narrow. It does not write "
        "the summary prose from scratch; it selects IDs. Code then reconstructs "
        "values, units, ranges, pages, excerpts, repeat groups, and mandatory "
        "uncertainties from the validated analysis. That is why the summary cannot "
        "invent a number. Repeat grouping is deterministic: normalized analyte, "
        "matching units, patient measurements only, and no invented chronology. "
        "In the UI, the compact view stays readable while every full evidence "
        "record remains one expander away."
    )
    return slide


# ---------------------------------------------------------------------------
# Slide 10 - Presenter 4
# ---------------------------------------------------------------------------


def slide_p4(prs):
    slide = base_slide(prs)
    header(
        slide,
        "Presenter 4  \u00b7  recommendation + approval",
        "Recommendations that cite \u2014 or decline",
        "Every suggestion needs both patient evidence and guideline evidence. "
        "If the guidance cannot support one, the system abstains.",
    )

    # Inputs
    rect(slide, MX, 1.64, 2.95, 2.2, fill=CARD_ALT, line=BORDER, radius=0.06)
    add_text(slide, MX + 0.18, 1.76, 2.55, 0.25, [p("Validated inputs", sz=11, b=True, c=TEXT, after=0)])
    inputs = ["Patient findings + evidence", "Validated summary", "Bounded guideline passages"]
    y = 2.12
    for label in inputs:
        chip(slide, MX + 0.18, y, 2.58, 0.42, label, sz=8.2, color=MUTED)
        y += 0.54

    # Agent
    rect(slide, 3.76, 1.64, 3.35, 2.2, fill=RGBColor(0x10, 0x27, 0x2B), line=RGBColor(0x1F, 0x5A, 0x55), radius=0.06)
    add_text(slide, 3.94, 1.8, 2.95, 0.55, [p("Recommendation Agent", sz=12.5, b=True, c=TEAL, after=0, line=1.0)])
    add_text(
        slide,
        3.94,
        2.42,
        2.95,
        1.0,
        [p("Draft review suggestions with both evidence types \u2014 or decline.", sz=9, c=MUTED, after=0, line=1.15)],
    )
    add_text(slide, 3.94, 3.5, 2.95, 0.25, [p("Gemini \u00b7 bounded structured call", sz=7.6, c=AMBER, after=0)])

    # Outcomes
    rect(slide, 7.44, 1.64, 5.34, 1.0, fill=CARD, line=BORDER, radius=0.06)
    add_text(slide, 7.62, 1.72, 4.9, 0.24, [p("Suggestion", sz=10.5, b=True, c=GREEN, after=0)])
    add_text(slide, 7.62, 1.98, 4.95, 0.26, [p("Text + finding IDs + passage IDs + exact quotes", sz=8.4, c=MUTED, after=0)])
    add_text(slide, 7.62, 2.28, 4.95, 0.25, [p("Status: Draft / Pending review", sz=8.2, b=True, c=AMBER, after=0)])

    rect(slide, 7.44, 2.74, 5.34, 0.95, fill=CARD, line=BORDER, radius=0.06)
    add_text(slide, 7.62, 2.84, 4.9, 0.24, [p("Abstention", sz=10.5, b=True, c=RED, after=0)])
    add_text(slide, 7.62, 3.1, 4.95, 0.5, [p("absent  \u00b7  unrelated  \u00b7  contradictory  \u00b7  insufficient guidance", sz=8.5, c=MUTED, after=0, line=1.05)])

    right_arrow_shape(slide, 3.47, 2.6, 0.2, 0.16)
    right_arrow_shape(slide, 7.18, 2.6, 0.2, 0.16)
    down_arrow(slide, 10.1, 3.74, 0.18, color=AMBER, w=0.14)

    # Human review band
    rect(slide, MX, 3.95, CW, 1.5, fill=RGBColor(0x2A, 0x20, 0x0E), line=RGBColor(0x6B, 0x4F, 0x18), radius=0.05)
    add_text(slide, MX + 0.2, 4.05, 6.0, 0.24, [p("HUMAN REVIEW  \u00b7  binds to one exact draft", sz=8, b=True, c=AMBER, after=0)])
    review = [
        ("Exact draft", "immutable content hash", "content-addressed; any edit changes it"),
        ("Approve / reject", "explicit action + notes", "no automatic approval path"),
        ("Decision record", "bound to that draft hash", "reloadable audit trail"),
        ("Edit", "creates a new pending version", "approval is never inherited"),
    ]
    x = MX + 0.2
    for idx, (title, sub, detail) in enumerate(review):
        add_text(slide, x, 4.42, 0.3, 0.25, [p(str(idx + 1), sz=9, b=True, c=AMBER, after=0)])
        add_text(slide, x + 0.3, 4.4, 2.55, 0.25, [p(title, sz=9.8, b=True, c=TEXT, after=0)])
        add_text(slide, x + 0.3, 4.68, 2.55, 0.25, [p(sub, sz=8.2, c=RGBColor(0xD6, 0xB9, 0x7A), after=0, line=1.03)])
        add_text(slide, x + 0.3, 4.98, 2.55, 0.25, [p(detail, sz=7.6, c=DIM, after=0, line=1.0)])
        x += 3.0

    # Hard limits
    rect(slide, MX, 5.65, CW, 0.5, fill=CARD, line=BORDER, radius=0.2)
    add_text(
        slide,
        MX + 0.22,
        5.75,
        CW - 0.44,
        0.3,
        [
            rp(
                [
                    ("Hard limits:  ", {"sz": 8.6, "b": True, "c": RED}),
                    (
                        ("cannot diagnose  \u00b7  cannot prescribe or change medication  "
                        "\u00b7  cannot provide doses  \u00b7  cannot assign urgency"),
                        {"sz": 8.6, "c": TEXT},
                    ),
                ],
                after=0,
            )
        ],
    )

    add_text(
        slide,
        MX,
        6.32,
        CW,
        0.28,
        [
            p(
                "Approval records a review decision. It is not clinical certification, "
                "and editing never inherits a prior approval.",
                sz=8.4,
                c=DIM,
                after=0,
            )
        ],
    )

    slide.notes_slide.notes_text_frame.text = (
        "Presenter 4. The recommendation is the most constrained agent. It receives "
        "only validated findings, the validated summary, and a bounded set of "
        "retrieved passages. Each suggestion must quote both patient evidence and "
        "guideline evidence, and those quotes are validated against the sources. "
        "If the passages are absent, unrelated, contradictory, or insufficient, "
        "the correct output is an abstention. Then the human step: approval binds "
        "to the exact draft hash. Editing creates a new pending version, so silence "
        "or a previous approval can never leak into a changed draft."
    )
    return slide


# ---------------------------------------------------------------------------
# Slide 11 - Presenter 5
# ---------------------------------------------------------------------------


def slide_p5(prs):
    slide = base_slide(prs)
    header(
        slide,
        "Presenter 5  \u00b7  orchestration + persistence",
        "Orchestration, parallelism, durable reuse",
        "Connect the agents without hiding failures \u2014 or repeating work that "
        "has already been validated.",
    )

    # Left: compact orchestration diagram
    rect(slide, MX, 1.64, 6.08, 4.6, fill=CARD_ALT, line=BORDER, radius=0.05)
    add_text(slide, MX + 0.2, 1.76, 5.4, 0.25, [p("CrewAI Flow", sz=12.5, b=True, c=BLUE, after=0)])
    add_text(slide, MX + 0.2, 2.04, 5.6, 0.22, [p("analysis first, then one parallel fork \u00b7 max 2 workers", sz=8.3, c=MUTED, after=0)])

    rect(slide, 1.55, 2.38, 3.95, 0.5, fill=RGBColor(0x10, 0x27, 0x2B), line=RGBColor(0x1F, 0x5A, 0x55), radius=0.15)
    add_text(slide, 1.68, 2.47, 3.7, 0.3, [p("1  Analysis resolved", sz=9.5, b=True, c=TEAL, after=0)])

    down_arrow(slide, 2.15, 2.93, 0.16, color=BLUE, w=0.13)
    down_arrow(slide, 4.9, 2.93, 0.16, color=BLUE, w=0.13)

    rect(slide, 0.85, 3.12, 2.55, 0.58, fill=CARD, line=BORDER, radius=0.12)
    add_text(slide, 1.0, 3.22, 2.3, 0.3, [p("2a  Summary Agent", sz=9.3, b=True, c=TEAL, after=0)])
    rect(slide, 3.65, 3.12, 2.55, 0.58, fill=RGBColor(0x1A, 0x16, 0x2E), line=RGBColor(0x4A, 0x3C, 0x78), radius=0.12)
    add_text(slide, 3.8, 3.22, 2.3, 0.3, [p("2b  Guideline retrieval", sz=9.3, b=True, c=PURPLE, after=0)])

    down_arrow(slide, 2.15, 3.76, 0.16, color=BLUE, w=0.13)
    down_arrow(slide, 4.9, 3.76, 0.16, color=BLUE, w=0.13)

    rect(slide, 1.55, 3.96, 3.95, 0.46, fill=CARD, line=BORDER, radius=0.15)
    add_text(slide, 1.68, 4.05, 3.7, 0.28, [p("3  Join + fingerprint re-check", sz=9.3, b=True, c=TEXT, after=0)])

    down_arrow(slide, 3.52, 4.46, 0.15, color=AMBER, w=0.13)
    rect(slide, 1.55, 4.65, 3.95, 0.46, fill=RGBColor(0x2A, 0x20, 0x0E), line=RGBColor(0x6B, 0x4F, 0x18), radius=0.15)
    add_text(slide, 1.68, 4.74, 3.7, 0.28, [p("4  Recommendation draft", sz=9.3, b=True, c=AMBER, after=0)])

    down_arrow(slide, 3.52, 5.15, 0.15, color=AMBER, w=0.13)
    rect(slide, 1.55, 5.34, 3.95, 0.46, fill=CARD, line=BORDER, radius=0.15)
    add_text(slide, 1.68, 5.43, 3.7, 0.28, [p("5  Review route", sz=9.3, b=True, c=TEXT, after=0)])

    add_text(
        slide,
        0.72,
        5.92,
        5.7,
        0.28,
        [p("Workers own no shared SQLite or Chroma connection.", sz=7.8, c=DIM, after=0)],
    )

    # Right: persistence
    rect(slide, 6.8, 1.64, 5.98, 4.6, fill=CARD, line=BORDER, radius=0.05)
    add_text(slide, 7.0, 1.76, 5.4, 0.25, [p("SQLite: exact, immutable reuse", sz=12.5, b=True, c=GREEN, after=0)])

    facts = [
        ("Immutable records", "analysis, summary, draft, and review decision"),
        ("Compatibility keys", "exact upstream hashes + settings + prompt/schema + runtime versions"),
        ("Guideline binding", "retrieval settings + document hashes + full index fingerprint"),
        ("Validated loads", "every load passes Pydantic and evidence re-checks"),
        ("Failure isolation", "saved stages survive; failures are sanitized and routed"),
    ]
    y = 2.14
    for title, body in facts:
        add_text(slide, 7.0, y, 0.18, 0.22, [p("\u25aa", sz=8, c=GREEN, after=0)])
        add_text(slide, 7.22, y - 0.01, 5.25, 0.22, [p(title, sz=9.6, b=True, c=TEXT, after=0)])
        add_text(slide, 7.22, y + 0.22, 5.25, 0.36, [p(body, sz=8.2, c=MUTED, after=0, line=1.03)])
        y += 0.72

    rect(slide, 7.0, 5.58, 5.5, 0.46, fill=RGBColor(0x0D, 0x2A, 0x2C), line=RGBColor(0x1F, 0x5A, 0x55), radius=0.18)
    add_text(
        slide,
        7.16,
        5.68,
        5.2,
        0.28,
        [p("Compatible result \u2192 zero new Gemini calls", sz=9, b=True, c=TEAL, after=0)],
    )

    # Route strip
    add_text(slide, MX, 6.4, 2.4, 0.2, [p("AUDITABLE ROUTES", sz=7.8, b=True, c=DIM, after=0)])
    routes = ["reuse", "generate", "retrieve", "abstain", "await_review", "approved", "rejected", "failed"]
    x = MX
    for label in routes:
        chip(slide, x, 6.63, 1.42, 0.26, label, sz=7.5, color=MUTED)
        x += 1.5

    slide.notes_slide.notes_text_frame.text = (
        "Presenter 5. The orchestration claim is precise: analysis resolves first; "
        "then one AI branch and one local retrieval branch run concurrently with "
        "two workers. It is not multiple AI agents in parallel. The join waits for "
        "both, then the flow re-checks the guideline fingerprint before building "
        "or reusing a recommendation. On the right, persistence is what makes the "
        "demo efficient without making it unsafe: immutable records, compatibility "
        "keys that bind exact upstream hashes and settings, and full revalidation "
        "on load. A failure preserves any stage that already validated."
    )
    return slide


# ---------------------------------------------------------------------------
# Slide 12 - Verification + limits
# ---------------------------------------------------------------------------


def slide_close(prs):
    slide = base_slide(prs)
    header(
        slide,
        "Verification + boundaries",
        "Built for trust \u2014 honest about limits",
        "The project is tested offline, demonstrated without live services, and "
        "explicit about what it does not prove.",
    )

    # Verification card
    rect(slide, MX, 1.64, 5.95, 3.88, fill=CARD, line=BORDER, radius=0.05)
    add_text(slide, MX + 0.22, 1.76, 5.4, 0.25, [p("Verification", sz=12.5, b=True, c=TEAL, after=0)])

    add_text(slide, MX + 0.22, 2.12, 1.55, 0.72, [p("173", sz=38, b=True, c=TEAL, after=0, line=0.9)])
    add_text(slide, MX + 1.55, 2.24, 1.6, 0.28, [p("tests passed", sz=10, b=True, c=TEXT, after=0)])
    add_text(slide, MX + 1.55, 2.5, 3.9, 0.28, [p("1 optional real-embedding test skipped", sz=8.2, c=MUTED, after=0)])

    checks = [
        "ruff check . clean; lockfile and sync checks pass",
        "Mocked end-to-end demo: no API key, no network, no model download",
        "Synthetic report plus three synthetic guideline PDFs",
        "Live Gemini smoke test: explicit --live, one request, retries disabled",
        "Validation re-runs on every saved-result load",
    ]
    y = 3.0
    for item in checks:
        add_text(slide, MX + 0.22, y, 0.2, 0.22, [p("\u2713", sz=9, b=True, c=GREEN, after=0)])
        add_text(slide, MX + 0.48, y - 0.01, 5.2, 0.32, [p(item, sz=8.7, c=MUTED, after=0, line=1.0)])
        y += 0.39

    add_text(
        slide,
        MX + 0.22,
        5.05,
        5.5,
        0.3,
        [
            rp(
                [
                    ("Demo:  ", {"sz": 8.2, "b": True, "c": TEAL}),
                    (
                        "upload \u2192 run workflow \u2192 inspect evidence \u2192 approve / reject",
                        {"sz": 8.2, "c": MUTED},
                    ),
                ],
                after=0,
            )
        ],
    )

    # Boundaries card
    rect(slide, 6.68, 1.64, 6.1, 3.88, fill=RGBColor(0x2A, 0x20, 0x0E), line=RGBColor(0x6B, 0x4F, 0x18), radius=0.05)
    add_text(slide, 6.88, 1.76, 5.5, 0.25, [p("Boundaries we state clearly", sz=12.5, b=True, c=AMBER, after=0)])
    limits = [
        "Educational capstone on synthetic data; not a medical device.",
        "Selectable-text PDFs only; OCR is out of scope.",
        "Range parsing is deliberately narrow; ambiguity goes to review.",
        "Similarity is retrieval ranking, not clinical confidence.",
        "Local SQLite is not a multi-user isolation boundary.",
        "Source validation proves text exists, not that guidance is correct.",
    ]
    add_text(
        slide,
        6.88,
        2.16,
        5.55,
        2.0,
        [bullet(item, sz=8.9, bc=AMBER, after=8) for item in limits],
    )
    rect(slide, 6.88, 4.35, 5.6, 0.62, fill=RGBColor(0x1A, 0x16, 0x2E), line=RGBColor(0x4A, 0x3C, 0x78), radius=0.15)
    add_text(
        slide,
        7.05,
        4.45,
        5.3,
        0.45,
        [
            rp(
                [
                    ("No automatic AI call:  ", {"sz": 8.2, "b": True, "c": PURPLE}),
                    (
                        "upload, rerun, and restart stay local until an explicit action.",
                        {"sz": 8.2, "c": MUTED},
                    ),
                ],
                after=0,
                line=1.0,
            )
        ],
    )

    # Closing band
    rect(slide, MX, 5.72, CW, 0.98, fill=RGBColor(0x10, 0x27, 0x2B), line=RGBColor(0x1F, 0x5A, 0x55), radius=0.06)
    add_text(slide, MX + 0.25, 5.83, 1.3, 0.22, [p("WHAT WE BUILT", sz=7.8, b=True, c=TEAL, after=0)])
    add_text(
        slide,
        MX + 0.25,
        6.08,
        9.9,
        0.5,
        [
            p(
                "A transparent review assistant that shows its sources, declines "
                "when unsure, and leaves the decision to a human.",
                sz=12.5,
                b=True,
                c=TEXT,
                after=0,
                line=1.0,
            )
        ],
    )
    chip(
        slide,
        11.08,
        6.0,
        1.55,
        0.4,
        "Questions",
        fill=RGBColor(0x0D, 0x2A, 0x2C),
        line=RGBColor(0x1F, 0x5A, 0x55),
        color=TEAL,
        sz=9,
        bold=True,
    )

    slide.notes_slide.notes_text_frame.text = (
        "Presenter 4 closes, then the whole team takes questions. State the "
        "evidence: 173 tests pass, one optional real-embedding test is skipped by "
        "default, and the offline demo runs without an API key or network. Then "
        "read the boundaries aloud: this is a synthetic educational capstone, not "
        "a medical device; OCR and broad range parsing are out of scope; "
        "similarity is not confidence. Finish on the closing line: transparency, "
        "abstention, and human authority are the product."
    )
    return slide


# ---------------------------------------------------------------------------
# Build
# ---------------------------------------------------------------------------


def main():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)

    slide_title(prs)
    slide_problem(prs)
    slide_solution(prs)
    slide_team(prs)
    slide_architecture(prs)
    slide_workflow(prs)
    slide_p1(prs)
    slide_p2(prs)
    slide_p3(prs)
    slide_p4(prs)
    slide_p5(prs)
    slide_close(prs)

    for index, slide in enumerate(prs.slides, start=1):
        if index > 1:
            footer(slide, index)

    props = prs.core_properties
    props.title = "Clinical Report Assistant - Capstone Presentation"
    props.subject = "Source-linked clinical report review with bounded AI and human approval"
    props.author = "Capstone Team"
    props.comments = "Synthetic-data educational capstone. Not a medical device."

    output = Path(__file__).resolve().parent.parent / "Clinical_Report_Assistant_Capstone.pptx"
    prs.save(output)
    print(f"Saved {output}")


if __name__ == "__main__":
    main()
